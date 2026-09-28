import os
import sys

# Configure headless matplotlib backend for web server environment
os.environ["MPLBACKEND"] = "Agg"
os.environ["TK_SILENCE_DEPRECATION"] = "1"

# Prevent Tcl/Tk from crashing Flask worker threads during cross-thread garbage collection
try:
    import tkinter
    if hasattr(tkinter, "Variable"):
        tkinter.Variable.__del__ = lambda self: None
    if hasattr(tkinter, "Image"):
        tkinter.Image.__del__ = lambda self: None
except Exception:
    pass

import matplotlib
matplotlib.use("Agg")

import tempfile
import io
import math
import random
from pathlib import Path

from flask import Flask, render_template, jsonify, send_file, request, send_from_directory
from openpyxl import load_workbook

# Base directory and script directories
BASE_DIR = Path(__file__).resolve().parent
CHEM_DIR = BASE_DIR / "Chemical Python Script"
DST_DIR = BASE_DIR / "DST Python Script"
FSI_DIR = BASE_DIR / "FSI Python Script"
HYD_DIR = BASE_DIR / "HYD Python Script"
LLPL_DIR = BASE_DIR / "LL & PL Python Script"
PARAM_DIR = BASE_DIR / "Parameter Python Script"
SHRINK_DIR = BASE_DIR / "Shrinkage Limit Python Script"
TRIAX_DIR = BASE_DIR / "Triaxial Python Script"
UCS_DIR = BASE_DIR / "UCS Python Script"

for p in [str(CHEM_DIR), str(DST_DIR), str(FSI_DIR), str(HYD_DIR), str(LLPL_DIR), str(PARAM_DIR), str(SHRINK_DIR), str(TRIAX_DIR), str(UCS_DIR)]:
    if p not in sys.path:
        sys.path.insert(0, p)

# Direct module imports from original untouched Python scripts
import soil_models
import dst_generator
import excel_export_multi
import FSI_Data_Generator as fsi_engine
import LL_PL_Generator_15_Row_GUI_SemiLog_FlowCurve_FINAL_v3 as llpl_engine
import chemical_test_generator_v13_input_controlled_groundwater as chem_engine

app = Flask(__name__)

# Catalog of all tests
TESTS_CATALOG = [
    {
        "id": "chemical",
        "name": "Soil & Groundwater Chemical Test",
        "short_name": "Chemical",
        "category": "Chemical Properties",
        "standard": "IS 2720 (Part 22, 26, 27) / IS 3025",
        "description": "Multi-borehole synthetic soil and groundwater chemical testing analysis: pH, water-soluble chloride, water-soluble sulphate, and organic matter content with depth zones and groundwater proximity modeling.",
        "folder": "Chemical Python Script",
        "icon": "fas fa-flask-vial",
        "color": "#0d9488",
        "tags": ["pH", "Chloride", "Sulphate", "Organic Matter", "Groundwater", "IS 2720", "IS 3025"],
        "status": "Active"
    },
    {
        "id": "dst",
        "name": "Direct Shear Test (DST)",
        "short_name": "DST",
        "category": "Shear Strength",
        "standard": "IS 2720 (Part 13) / ASTM D3080",
        "description": "Multi-depth synthetic direct shear test data generator, normal stress envelopes, shear stress vs. horizontal displacement curves, target c/phi determination, and multi-specimen Excel export.",
        "folder": "DST Python Script",
        "icon": "fas fa-layer-group",
        "color": "#2563eb",
        "tags": ["Shear Strength", "Multi-Depth", "Cohesion & Phi", "Excel Export"],
        "status": "Active"
    },
    {
        "id": "fsi",
        "name": "Free Swell Index (FSI)",
        "short_name": "FSI",
        "category": "Expansive Soil Properties",
        "standard": "IS 2720 (Part 40)",
        "description": "Calculation of free swell index using kerosene and distilled water graduated cylinders, swell degree classification, and automated reporting.",
        "folder": "FSI Python Script",
        "icon": "fas fa-arrows-alt-v",
        "color": "#059669",
        "tags": ["Swelling", "Clay Soils", "Index Property", "Expansiveness"],
        "status": "Active"
    },
    {
        "id": "hyd",
        "name": "Hydrometer Analysis (HYD)",
        "short_name": "HYD",
        "category": "Grain Size Analysis",
        "standard": "IS 2720 (Part 4) / ASTM D422",
        "description": "Sedimentation analysis for fine-grained soils (< 75 microns), Stokes' law particle diameter determination, meniscus/temperature corrections, and semi-log gradation plotting.",
        "folder": "HYD Python Script",
        "icon": "fas fa-vial",
        "color": "#7c3aed",
        "tags": ["Gradation", "Stokes' Law", "Fine Grained", "Particle Size"],
        "status": "Active"
    },
    {
        "id": "ll_pl",
        "name": "Liquid Limit & Plastic Limit (LL & PL)",
        "short_name": "LL & PL",
        "category": "Atterberg Limits",
        "standard": "IS 2720 (Part 5) / ASTM D4318",
        "description": "Casagrande cup / cone penetrometer liquid limit, plastic limit thread rolling, plasticity index (PI), liquidity index, flow curve generation, and A-line soil plasticity chart.",
        "folder": "LL & PL Python Script",
        "icon": "fas fa-water",
        "color": "#0284c7",
        "tags": ["Atterberg Limits", "Plasticity Index", "A-Line Chart", "USCS Classification"],
        "status": "Active"
    },
    {
        "id": "parameter",
        "name": "Engineering Strength Parameter Generator",
        "short_name": "Parameters",
        "category": "Strength Parameters",
        "standard": "IS 2131 / IS 2720",
        "description": "Multi-sheet synthetic engineering strength parameter generator with IS 2131 SPT N-value overburden and dilatancy corrections, PI-dependent undrained cohesion (Su/Cu/UCS), and soil-specific drained c'/phi'.",
        "folder": "Parameter Python Script",
        "icon": "fas fa-calculator",
        "color": "#d97706",
        "tags": ["IS 2131", "SPT Correction", "Su & UCS", "Drained c & phi", "Multi-Sheet"],
        "status": "Active"
    },
    {
        "id": "shrinkage",
        "name": "Shrinkage Limit Test (SL)",
        "short_name": "Shrinkage",
        "category": "Atterberg Limits",
        "standard": "IS 2720 (Part 6) / ASTM D4943",
        "description": "Determination of shrinkage limit, shrinkage ratio, volumetric shrinkage, and linear shrinkage using mercury displacement / wax coating method.",
        "folder": "Shrinkage Limit Python Script",
        "icon": "fas fa-compress-arrows-alt",
        "color": "#dc2626",
        "tags": ["Shrinkage Limit", "Volumetric Change", "Mercury Method", "Shrinkage Ratio"],
        "status": "Active"
    },
    {
        "id": "triaxial",
        "name": "Triaxial Shear Test",
        "short_name": "Triaxial",
        "category": "Shear Strength",
        "standard": "IS 2720 (Part 11 & 12) / ASTM D2850",
        "description": "Unconsolidated Undrained (UU), Consolidated Undrained (CU), and Consolidated Drained (CD) triaxial testing analysis, Mohr-Coulomb stress circles, failure envelopes, and pore pressure tracking.",
        "folder": "Triaxial Python Script",
        "icon": "fas fa-cube",
        "color": "#4f46e5",
        "tags": ["Mohr Circles", "Triaxial", "UU / CU / CD", "Pore Pressure"],
        "status": "Active"
    },
    {
        "id": "ucs",
        "name": "Unconfined Compressive Strength (UCS)",
        "short_name": "UCS",
        "category": "Compressive Strength",
        "standard": "IS 2720 (Part 10) / ASTM D2166",
        "description": "Stress-strain curve recording for cohesive soils, unconfined compressive strength (qu), undrained shear strength (cu = qu/2), sensitivity, and failure strain determination.",
        "folder": "UCS Python Script",
        "icon": "fas fa-tachometer-alt",
        "color": "#0891b2",
        "tags": ["Compressive Strength", "Stress-Strain", "Cohesive Soils", "Undrained Cohesion"],
        "status": "Active"
    }
]

@app.route("/")
def index():
    return render_template("index.html", tests=TESTS_CATALOG)

@app.route("/tests/<test_id>")
def test_page(test_id):
    test_info = next((t for t in TESTS_CATALOG if t["id"] == test_id), None)
    if not test_info:
        return "Test not found", 404
    
    template_name = f"tests/{test_id}.html"
    if os.path.exists(os.path.join(app.root_path, "templates", "tests", f"{test_id}.html")):
        return render_template(template_name, test=test_info)
    
    return render_template("test_placeholder.html", test=test_info)

@app.route("/outputs/<path:filepath>")
def serve_output_file(filepath):
    outputs_dir = BASE_DIR / "outputs"
    as_attachment = request.args.get("download", "").lower() in ("1", "true")
    return send_from_directory(outputs_dir, filepath, as_attachment=as_attachment)

@app.route("/api/tests")
def api_tests():
    return jsonify({"success": True, "tests": TESTS_CATALOG})

# =========================================================================
# Chemical (Soil & Groundwater Chemical Test Generator)
# =========================================================================

@app.route("/api/chemical/defaults", methods=["GET"])
def api_chemical_defaults():
    return jsonify({
        "success": True,
        "environments": list(chem_engine.SOIL_RANGES.keys()),
        "soil_types": list(chem_engine.SOIL_EFFECT.keys()),
        "condition_options": chem_engine.CONDITION_OPTIONS,
        "salinity_control_options": chem_engine.SALINITY_CONTROL_OPTIONS,
        "gw_condition_options": chem_engine.GW_CONDITION_OPTIONS,
        "gw_salinity_control_options": chem_engine.GW_SALINITY_CONTROL_OPTIONS,
        "residue_options": ["None", "Slight", "Visible", "Strong"]
    })

@app.route("/api/chemical/generate", methods=["POST"])
def api_chemical_generate():
    data = request.json or {}
    project = str(data.get("project", "")).strip()
    location = str(data.get("location", "")).strip()
    env = str(data.get("environment", "Normal River")).strip()
    if env not in chem_engine.SOIL_RANGES:
        env = "Normal River"

    seed_val = data.get("seed", "")
    try:
        seed = int(seed_val)
        if seed < 0:
            return jsonify({"error": "Seed must be non-negative."}), 400
    except (ValueError, TypeError):
        seed = random.randint(100000, 999999)

    random.seed(seed)
    boreholes = data.get("boreholes", [])
    if not boreholes:
        return jsonify({"error": "At least one borehole row is required."}), 400

    soil_results = []
    gw_results = []
    summary_rows = []

    # Validation loop
    for i, b in enumerate(boreholes, 1):
        bh_id = str(b.get("bh_id", f"BH-{i:02d}")).strip()
        if not bh_id:
            return jsonify({"error": f"Row {i}: BH ID is required."}), 400
        try:
            d1 = float(b.get("s1_depth", 1.5))
            d2 = float(b.get("s2_depth", 4.5))
            dg = float(b.get("gw_depth", 8.0))
            if d1 <= 0 or d2 <= 0 or dg <= 0:
                return jsonify({"error": f"{bh_id}: Depths must be greater than 0."}), 400
            if d1 >= d2:
                return jsonify({"error": f"{bh_id}: S-1 depth ({d1} m) must be less than S-2 depth ({d2} m)."}), 400
        except (ValueError, TypeError):
            return jsonify({"error": f"{bh_id}: Depths must be valid numbers."}), 400

        sc1 = str(b.get("s1_salinity", "Normal"))
        sc2 = str(b.get("s2_salinity", "Normal"))
        gwsc = str(b.get("gw_salinity", "Normal"))

        if sc1 not in chem_engine.SALINITY_CONTROL_OPTIONS:
            return jsonify({"error": f"{bh_id}: Invalid S-1 salinity control '{sc1}'."}), 400
        if sc2 not in chem_engine.SALINITY_CONTROL_OPTIONS:
            return jsonify({"error": f"{bh_id}: Invalid S-2 salinity control '{sc2}'."}), 400
        if gwsc not in chem_engine.GW_SALINITY_CONTROL_OPTIONS:
            return jsonify({"error": f"{bh_id}: Invalid groundwater salinity control '{gwsc}'."}), 400

    def validate_soil_result(env_name, cl_val, so_val, ph_val, res_val):
        r = chem_engine.SOIL_RANGES[env_name]
        if not (r["pH"]["Low"][0] <= ph_val <= r["pH"]["Very High"][1]): return "CHECK: pH"
        if cl_val <= 0 or so_val <= 0: return "CHECK: salt"
        if res_val == "None" and env_name != "Desert River" and (cl_val > 2200 or so_val > 2700): return "CHECK: residue/salt"
        return "PASS"

    def validate_gw_result(env_name, cl_val, so_val, ph_val):
        r = chem_engine.GW_RANGES[env_name]
        if not (r["pH"]["Low"][0] <= ph_val <= r["pH"]["Very High"][1]): return "CHECK: pH"
        if cl_val <= 0 or so_val <= 0: return "CHECK: salt"
        return "PASS"

    # Generation loop
    for i, b in enumerate(boreholes, 1):
        bh_id = str(b.get("bh_id", f"BH-{i:02d}")).strip()
        d1 = float(b.get("s1_depth", 1.5))
        t1 = str(b.get("s1_type", "CL"))
        c1 = str(b.get("s1_condition", "Moist"))
        sc1 = str(b.get("s1_salinity", "Normal"))

        d2 = float(b.get("s2_depth", 4.5))
        t2 = str(b.get("s2_type", "CL"))
        c2 = str(b.get("s2_condition", "Moist"))
        sc2 = str(b.get("s2_salinity", "Normal"))

        dg = float(b.get("gw_depth", 8.0))
        gwc = str(b.get("gw_condition", "Normal"))
        gwsc = str(b.get("gw_salinity", "Normal"))
        residue = str(b.get("white_residue", "None"))

        profile = chem_engine.salt_condition(env, residue)
        dominance = chem_engine.choose_dominance(env, residue)
        ph_anchor = chem_engine.correlated_natural_ground_ph() if env == "Natural Ground / Bare Land" else None

        s1 = chem_engine.soil_result(env, t1, residue, profile, 1, c1, d1, dg, dominance, sc1, ph_anchor)
        s2 = chem_engine.soil_result(env, t2, residue, profile, 2, c2, d2, dg, dominance, sc2, ph_anchor)
        gw = chem_engine.groundwater_result(env, profile, gwc, dominance, gwsc)

        for idx, (depth, typ, cond_in, res_data) in enumerate([(d1, t1, c1, s1), (d2, t2, c2, s2)], 1):
            p, cl, so, om, cond_out, zone, dom = res_data
            valid = validate_soil_result(env, cl, so, p, residue)
            soil_results.append({
                "bh_id": bh_id,
                "sample": f"S-{idx}",
                "depth": depth,
                "zone": zone,
                "soil_type": typ,
                "condition": cond_in,
                "salinity_control": cond_out,
                "environment": env,
                "ph": p,
                "chloride": cl,
                "sulphate": so,
                "organic_matter": om,
                "white_residue": residue,
                "chemical_dominance": dom,
                "validation": valid
            })

        validgw = validate_gw_result(env, gw[1], gw[2], gw[0])
        gw_results.append({
            "bh_id": bh_id,
            "sample": "GW-1",
            "depth": dg,
            "environment": env,
            "condition": gwc,
            "ph": gw[0],
            "chloride": gw[1],
            "sulphate": gw[2],
            "salinity_condition": gw[3],
            "chemical_dominance": gw[4],
            "validation": validgw
        })

        summary_rows.append({
            "bh_id": bh_id,
            "environment": env,
            "s1_depth": d1,
            "s1_zone": chem_engine.depth_zone(d1),
            "s1_type": t1,
            "s1_condition": c1,
            "s1_salinity": sc1,
            "s2_depth": d2,
            "s2_zone": chem_engine.depth_zone(d2),
            "s2_type": t2,
            "s2_condition": c2,
            "s2_salinity": sc2,
            "gw_depth": dg,
            "gw_condition": gwc,
            "gw_salinity": gwsc,
            "white_residue": residue,
            "chemical_dominance": dominance,
            "s1_cl": s1[1],
            "s2_cl": s2[1],
            "gw_cl": gw[1],
            "s1_so4": s1[2],
            "s2_so4": s2[2],
            "gw_so4": gw[2],
            "overall_check": "PASS"
        })

    method_rows = [
        {"parameter": "pH", "sample": "Soil", "standard": "IS 2720 (Part 26)", "method": "Electrometric / glass electrode", "unit": "pH"},
        {"parameter": "Chloride", "sample": "Soil", "standard": "IS 2720 (Part 27)", "method": "Water extract; argentometric titration", "unit": "mg/kg"},
        {"parameter": "Sulphate", "sample": "Soil", "standard": "IS 2720 (Part 27)", "method": "Water extract; gravimetric / turbidimetric", "unit": "mg/kg"},
        {"parameter": "Organic Matter", "sample": "Soil", "standard": "IS 2720 (Part 22)", "method": "Wet oxidation", "unit": "%"},
        {"parameter": "pH", "sample": "Groundwater", "standard": "IS 3025 (Part 11)", "method": "Electrometric / glass electrode", "unit": "pH"},
        {"parameter": "Chloride", "sample": "Groundwater", "standard": "IS 3025 (Part 32)", "method": "Argentometric titration", "unit": "mg/L"},
        {"parameter": "Sulphate", "sample": "Groundwater", "standard": "IS 3025 (Part 24)", "method": "Turbidimetric / gravimetric", "unit": "mg/L"}
    ]

    return jsonify({
        "success": True,
        "seed": seed,
        "environment": env,
        "project": project,
        "location": location,
        "count": len(boreholes),
        "soil_results": soil_results,
        "gw_results": gw_results,
        "summary_rows": summary_rows,
        "methods": method_rows,
        "status_text": f"Generated {len(boreholes)} boreholes: {len(soil_results)} soil tests + {len(gw_results)} groundwater tests. Seed={seed}"
    })

@app.route("/api/chemical/export_excel", methods=["POST"])
def api_chemical_export_excel():
    data = request.json or {}
    soil_results = data.get("soil_results", [])
    gw_results = data.get("gw_results", [])
    summary_rows = data.get("summary_rows", [])
    project = str(data.get("project", "")).strip()
    location = str(data.get("location", "")).strip()
    env = str(data.get("environment", "Normal River")).strip()
    seed = str(data.get("seed", "")).strip()

    if not soil_results:
        return jsonify({"error": "No chemical test results to export. Generate data first."}), 400

    from openpyxl import Workbook
    from openpyxl.styles import Font, Alignment
    from openpyxl.utils import get_column_letter

    wb = Workbook()
    ws = wb.active
    ws.title = "Soil Chemical Results"

    soil_table_rows = []
    for r in soil_results:
        soil_table_rows.append([
            r.get("bh_id", ""), r.get("sample", ""), r.get("depth", ""), r.get("zone", ""),
            r.get("soil_type", ""), r.get("condition", ""), r.get("salinity_control", ""),
            r.get("environment", ""), r.get("ph", ""), r.get("chloride", ""),
            r.get("sulphate", ""), r.get("organic_matter", ""), r.get("white_residue", ""),
            r.get("chemical_dominance", ""), r.get("validation", "")
        ])

    gw_table_rows = []
    for r in gw_results:
        gw_table_rows.append([
            r.get("bh_id", ""), r.get("sample", ""), r.get("depth", ""), r.get("environment", ""),
            r.get("condition", ""), r.get("ph", ""), r.get("chloride", ""),
            r.get("sulphate", ""), r.get("salinity_condition", ""),
            r.get("chemical_dominance", ""), r.get("validation", "")
        ])

    summary_table_rows = []
    for r in summary_rows:
        summary_table_rows.append([
            r.get("bh_id", ""), r.get("environment", ""), r.get("s1_depth", ""), r.get("s1_zone", ""),
            r.get("s1_type", ""), r.get("s1_condition", ""), r.get("s1_salinity", ""),
            r.get("s2_depth", ""), r.get("s2_zone", ""), r.get("s2_type", ""),
            r.get("s2_condition", ""), r.get("s2_salinity", ""), r.get("gw_depth", ""),
            r.get("gw_condition", ""), r.get("gw_salinity", ""), r.get("white_residue", ""),
            r.get("chemical_dominance", ""), r.get("s1_cl", ""), r.get("s2_cl", ""),
            r.get("gw_cl", ""), r.get("s1_so4", ""), r.get("s2_so4", ""),
            r.get("gw_so4", ""), r.get("overall_check", "PASS")
        ])

    methods = [
        ["Parameter", "Sample", "IS Code / Standard", "Method", "Unit"],
        ["pH", "Soil", "IS 2720 (Part 26)", "Electrometric / glass electrode", "pH"],
        ["Chloride", "Soil", "IS 2720 (Part 27)", "Water extract; argentometric titration", "mg/kg"],
        ["Sulphate", "Soil", "IS 2720 (Part 27)", "Water extract; gravimetric / turbidimetric", "mg/kg"],
        ["Organic Matter", "Soil", "IS 2720 (Part 22)", "Wet oxidation", "%"],
        ["pH", "Groundwater", "IS 3025 (Part 11)", "Electrometric / glass electrode", "pH"],
        ["Chloride", "Groundwater", "IS 3025 (Part 32)", "Argentometric titration", "mg/L"],
        ["Sulphate", "Groundwater", "IS 3025 (Part 24)", "Turbidimetric / gravimetric", "mg/L"],
    ]

    settings = [
        ["Project", project], ["Location", location],
        ["Environment", env], ["Boreholes", len(summary_rows)],
        ["Soil samples / borehole", 2], ["Groundwater samples / borehole", 1],
        ["Seed", seed], ["Depth Model", "Surface / Intermediate / Deep zones with groundwater proximity influence"],
        ["Import Template", "Use the Borehole Input worksheet for direct import."],
        ["Note", "Synthetic data for testing/workflow purposes; not actual laboratory measurements."]
    ]

    sheets_data = [
        ("Soil Chemical Results", soil_table_rows),
        ("Groundwater Chemical Results", gw_table_rows),
        ("Borehole Summary", summary_table_rows),
        ("Test Methods", methods),
        ("Generation Settings", settings)
    ]

    for idx, (name, s_data) in enumerate(sheets_data):
        if idx == 0:
            sh = ws
        else:
            sh = wb.create_sheet(name)

        if name == "Soil Chemical Results":
            headers = ["BH ID", "Sample", "Depth (m)", "Depth Zone", "Soil Type", "Condition", "Salinity Control", "Environment", "pH", "Chloride (mg/kg)", "Sulphate (mg/kg)", "Organic Matter (%)", "White Residue", "Chemical Dominance", "Validation"]
            full_data = [headers] + s_data
        elif name == "Groundwater Chemical Results":
            headers = ["BH ID", "Sample", "Depth (m)", "Environment", "Condition", "pH", "Chloride (mg/L)", "Sulphate (mg/L)", "Salinity Condition", "Chemical Dominance", "Validation"]
            full_data = [headers] + s_data
        elif name == "Borehole Summary":
            headers = ["BH ID", "Environment", "S-1 Depth", "S-1 Zone", "S-1 Type", "S-1 Condition", "S-1 Salinity Control", "S-2 Depth", "S-2 Zone", "S-2 Type", "S-2 Condition", "S-2 Salinity Control", "GW Depth", "GW Condition", "GW Salinity Control", "White Residue", "Chemical Dominance", "S-1 Cl", "S-2 Cl", "GW Cl", "S-1 SO4", "S-2 SO4", "GW SO4", "Overall Check"]
            full_data = [headers] + s_data
        else:
            full_data = s_data

        for r, row in enumerate(full_data, 1):
            for c, val in enumerate(row, 1):
                cell = sh.cell(r, c, val)
                cell.alignment = Alignment(horizontal="center", vertical="center")
                if r == 1:
                    cell.font = Font(bold=True)
        sh.freeze_panes = "A2"
        sh.auto_filter.ref = sh.dimensions
        for col in range(1, sh.max_column + 1):
            vals = [str(sh.cell(r, col).value or "") for r in range(1, min(sh.max_row, 100) + 1)]
            sh.column_dimensions[get_column_letter(col)].width = min(35, max(12, max(map(len, vals)) + 2))

    outputs_dir = BASE_DIR / "outputs" / "Chemical"
    outputs_dir.mkdir(parents=True, exist_ok=True)
    clean_proj = "".join(c for c in project if c.isalnum() or c in "-_") or "Chemical_Test_Report"
    filename = f"{clean_proj}_Results.xlsx"
    out_file = outputs_dir / filename
    wb.save(out_file)

    return jsonify({
        "success": True,
        "filename": filename,
        "excel_path": str(out_file),
        "download_url": f"/outputs/Chemical/{filename}"
    })

@app.route("/api/chemical/import_excel", methods=["POST"])
def api_chemical_import_excel():
    if "file" not in request.files:
        return jsonify({"error": "No file uploaded."}), 400
    file = request.files["file"]
    if not file.filename:
        return jsonify({"error": "Empty filename."}), 400

    required = [
        "BH ID", "S-1 Depth (m)", "S-1 Type", "S-1 Condition", "S-1 Salinity Control",
        "S-2 Depth (m)", "S-2 Type", "S-2 Condition", "S-2 Salinity Control",
        "GW Depth (m)", "GW Condition", "GW Salinity Control", "White Residue"
    ]

    try:
        wb = load_workbook(file, data_only=True)
        if "Borehole Input" in wb.sheetnames:
            ws = wb["Borehole Input"]
        else:
            ws = wb[wb.sheetnames[0]]

        rows = list(ws.iter_rows(values_only=True))
        if not rows:
            return jsonify({"error": "The selected workbook is empty."}), 400

        headers = [str(x).strip() if x is not None else "" for x in rows[0]]
        missing = [h for h in required if h not in headers]
        if missing:
            return jsonify({"error": "Required column(s) missing: " + ", ".join(missing)}), 400

        idx = {h: headers.index(h) for h in required}
        imported = []

        for excel_row, row in enumerate(rows[1:], start=2):
            if not any(v is not None and str(v).strip() != "" for v in row):
                continue

            vals = {}
            for h in required:
                value = row[idx[h]] if idx[h] < len(row) else None
                vals[h] = "" if value is None else str(value).strip()

            bh_id = vals["BH ID"]
            if not bh_id or bh_id.startswith("#"):
                continue

            imported.append({
                "bh_id": bh_id,
                "s1_depth": vals.get("S-1 Depth (m)", "1.5"),
                "s1_type": vals.get("S-1 Type", "CL"),
                "s1_condition": vals.get("S-1 Condition", "Moist"),
                "s1_salinity": vals.get("S-1 Salinity Control", "Normal"),
                "s2_depth": vals.get("S-2 Depth (m)", "4.5"),
                "s2_type": vals.get("S-2 Type", "CL"),
                "s2_condition": vals.get("S-2 Condition", "Moist"),
                "s2_salinity": vals.get("S-2 Salinity Control", "Normal"),
                "gw_depth": vals.get("GW Depth (m)", "8.0"),
                "gw_condition": vals.get("GW Condition", "Normal"),
                "gw_salinity": vals.get("GW Salinity Control", "Normal"),
                "white_residue": vals.get("White Residue", "None"),
            })

        if not imported:
            return jsonify({"error": "No valid borehole rows were found."}), 400

        settings = {}
        if "Generation Settings" in wb.sheetnames:
            sws = wb["Generation Settings"]
            for r in sws.iter_rows(min_row=1, max_col=2, values_only=True):
                if r[0] is not None:
                    settings[str(r[0]).strip()] = "" if r[1] is None else str(r[1]).strip()

        wb.close()

        return jsonify({
            "success": True,
            "count": len(imported),
            "boreholes": imported,
            "project": settings.get("Project", ""),
            "location": settings.get("Location", ""),
            "environment": settings.get("Environment", "Normal River"),
            "seed": settings.get("Seed", "")
        })
    except Exception as e:
        return jsonify({"error": f"Failed to import Excel: {str(e)}"}), 500

# =========================================================================
# DST (Direct Shear Test) Backend Endpoints - Exact Existing Logic
# =========================================================================

@app.route("/api/dst/soil_names", methods=["GET"])
def api_dst_soil_names():
    return jsonify({"soil_names": soil_models.get_soil_names()})

@app.route("/api/dst/classify_density", methods=["POST"])
def api_dst_classify_density():
    data = request.json or {}
    soil = data.get("soil", "")
    try:
        density = float(data.get("density", 0))
        behaviour = dst_generator.classify_density(soil, density)
    except (ValueError, TypeError):
        behaviour = "Enter density"
    return jsonify({"behaviour": behaviour})

@app.route("/api/dst/validate", methods=["POST"])
def api_dst_validate():
    try:
        data = request.json or {}
        rows = data.get("rows", [])
        warnings, errors = [], []

        bh_id = data.get("bh_id", "").strip()
        if not bh_id:
            errors.append("BH ID is required.")

        try:
            pr = float(data.get("pr_constant", 0))
            n1 = float(data.get("normal_1", 0))
            n2 = float(data.get("normal_2", 0))
            n3 = float(data.get("normal_3", 0))
            if pr <= 0: errors.append("PR Constant must be greater than zero.")
            if not (n1 > 0 and n2 > 0 and n3 > 0): errors.append("Normal stresses must be greater than zero.")
            if not (n1 < n2 < n3): errors.append("Normal stresses must satisfy Trial 1 < Trial 2 < Trial 3.")
        except (ValueError, TypeError):
            errors.append("PR Constant and all three normal stresses must be numeric.")

        if not rows:
            errors.append("At least one depth row is required.")

        for i, r in enumerate(rows, 1):
            try:
                depth_str = str(r.get("depth", "")).strip()
                density_str = str(r.get("density", "")).strip()
                c_str = str(r.get("c", "")).strip()
                phi_str = str(r.get("phi", "")).strip()

                if not depth_str or not density_str or not c_str or not phi_str:
                    errors.append(f"Row {i}: Depth, density, sample weight, c and phi must be numeric.")
                    continue

                depth = float(depth_str)
                density = float(density_str)
                c = float(c_str)
                phi = float(phi_str)
            except (ValueError, TypeError):
                errors.append(f"Row {i}: Depth, density, sample weight, c and phi must be numeric.")
                continue

            if r.get("sample_type") not in ("Undisturbed Sample", "Remolded"):
                errors.append(f"Row {i}: Select a valid sample type.")
            if depth < 0: errors.append(f"Row {i}: Sample depth cannot be negative.")
            if density <= 0: errors.append(f"Row {i}: Density must be greater than zero.")
            if c < 0: errors.append(f"Row {i}: Cohesion cannot be negative.")
            if phi < 0 or phi >= 90: errors.append(f"Row {i}: Friction angle must be between 0 and 90 degrees.")

            inside, msg = soil_models.validate_target_strength(r.get("soil", ""), c, phi)
            if not inside:
                warnings.append(f"Depth {depth:g} m:\n{msg}")

        return jsonify({"valid": len(errors) == 0, "errors": errors, "warnings": warnings, "count": len(rows)})
    except Exception as e:
        return jsonify({"valid": False, "errors": [str(e)], "warnings": []}), 400

@app.route("/api/dst/generate", methods=["POST"])
def api_dst_generate():
    try:
        data = request.json or {}
        rows = data.get("rows", [])
        bh_id = data.get("bh_id", "").strip()
        if not bh_id:
            return jsonify({"error": "BH ID is required."}), 400

        try:
            pr = float(data.get("pr_constant", 0.325))
            stresses = [float(data.get("normal_1", 0.5)), float(data.get("normal_2", 1.0)), float(data.get("normal_3", 1.5))]
        except (ValueError, TypeError):
            return jsonify({"error": "PR Constant and all normal stresses must be numeric."}), 400

        if not (stresses[0] < stresses[1] < stresses[2]):
            return jsonify({"error": "Normal stresses must satisfy Trial 1 < Trial 2 < Trial 3."}), 400

        if not rows:
            return jsonify({"error": "At least one depth row is required."}), 400

        seed_text = str(data.get("random_seed", "")).strip()
        base_seed = None if not seed_text or seed_text.upper() == "AUTO" else int(seed_text)

        results = []
        for i, item in enumerate(rows, 1):
            try:
                depth_val = float(item["depth"])
                density_val = float(item["density"])
                c_val = float(item["c"])
                phi_val = float(item["phi"])
                weight_val = float(item.get("weight") or (density_val * 90.0))
            except (ValueError, TypeError, KeyError):
                return jsonify({"error": f"Row {i}: Depth, density, sample weight, c and phi must be numeric."}), 400

            soil_type = item.get("soil", "Medium Dense Sand")
            sample_type = item.get("sample_type", "Undisturbed Sample")
            seed = None if base_seed is None else base_seed + (i - 1)

            trials = dst_generator.generate_three_trials(
                soil_type=soil_type, density=density_val, c=c_val, phi=phi_val,
                normal_stresses=stresses, seed=seed, n_points=61
            )
            summary = dst_generator.calculate_summary(trials)

            serialized_trials = []
            for t in trials:
                serialized_trials.append({
                    "normal_stress": float(t.normal_stress),
                    "displacement": [float(x) for x in t.displacement],
                    "shear_stress": [float(x) for x in t.shear_stress],
                    "failure_shear_stress": float(t.failure_shear_stress),
                    "peak_shear_stress": float(t.peak_shear_stress),
                    "peak_displacement": float(t.peak_displacement),
                    "residual_shear_stress": float(t.residual_shear_stress)
                })

            results.append({
                "depth": depth_val, "soil": soil_type, "sample_type": sample_type,
                "density": density_val, "weight": weight_val, "c": c_val, "phi": phi_val,
                "trials": serialized_trials,
                "summary": {
                    "calculated_c": float(summary["calculated_c"]),
                    "calculated_phi": float(summary["calculated_phi"]),
                    "r_squared": float(summary["r_squared"])
                }
            })

        return jsonify({"success": True, "results": results})
    except Exception as e:
        return jsonify({"error": str(e)}), 500

def _render_dst_plots(bh_id, results_raw, plot_dir):
    import matplotlib
    matplotlib.use("Agg")
    import matplotlib.pyplot as plt
    from matplotlib.patches import Rectangle
    from matplotlib.lines import Line2D
    import re

    def safe_name(text):
        text = str(text).strip()
        text = re.sub(r"[^A-Za-z0-9._-]+", "_", text)
        return text.strip("._") or "Result"

    plot_dir.mkdir(parents=True, exist_ok=True)
    # Clear existing old/stale plots in directory so previous runs don't linger
    for old_file in plot_dir.glob("*.png"):
        try:
            old_file.unlink()
        except Exception:
            pass
    created_files = []

    for result in results_raw:
        trials = result["trials"]
        summary = result["summary"]
        if isinstance(trials[0], dst_generator.DSTTrial):
            sigma = [float(t.normal_stress) for t in trials]
            tau = [float(t.peak_shear_stress) for t in trials]
        else:
            sigma = [float(t["normal_stress"]) for t in trials]
            tau = [float(t["peak_shear_stress"]) for t in trials]

        calc_c = float(summary["calculated_c"])
        calc_phi = float(summary["calculated_phi"])
        r2 = float(summary["r_squared"])
        slope = math.tan(math.radians(calc_phi))

        xmax = max(1.2, math.ceil(max(sigma) * 1.20 * 10) / 10)
        ymax = max(1.2, math.ceil(max(tau) * 1.20 * 10) / 10)
        x_step = 0.2 if xmax <= 1.2 else 0.5
        y_step = 0.2 if ymax <= 1.2 else 0.5

        fig = plt.figure(figsize=(15.36, 10.24), dpi=150, facecolor="white")
        fig.subplots_adjust(0, 0, 1, 1)
        fig.patches.append(Rectangle(
            (0.006, 0.006), 0.988, 0.988,
            transform=fig.transFigure, fill=False,
            linewidth=1.4, edgecolor="black"
        ))

        # Header
        fig.text(0.5, 0.965, "DIRECT SHEAR TEST – FAILURE ENVELOPE",
                 ha="center", va="center", fontsize=22,
                 fontweight="bold", color="#14255c")
        fig.text(0.5, 0.936, "IS 2720 (Part 13) – 1986",
                 ha="center", va="center", fontsize=12,
                 fontweight="bold", color="#14255c")

        # Top information table
        y_top, y_bottom = 0.842, 0.925
        x_edges = [0.008, 0.258, 0.508, 0.758, 0.992]
        for x in x_edges:
            fig.add_artist(Line2D([x, x], [y_top, y_bottom],
                                  transform=fig.transFigure, color="black", linewidth=0.9))
        fig.add_artist(Line2D([0.008, 0.992], [y_top, y_top],
                              transform=fig.transFigure, color="black", linewidth=0.9))
        fig.add_artist(Line2D([0.008, 0.992], [y_bottom, y_bottom],
                              transform=fig.transFigure, color="black", linewidth=0.9))

        sample_label = result["sample_type"]
        condition = "Remoulded" if sample_label == "Remolded" else "Undisturbed"
        info = [
            [("Borehole ID", bh_id), ("Depth", f"{float(result['depth']):.2f} m")],
            [("Sample Type", sample_label)],
            [("Shear Box Size", "60 mm × 60 mm")],
            [("Condition", condition)],
        ]
        for ci, entries in enumerate(info):
            xa, xb = x_edges[ci], x_edges[ci + 1]
            if ci == 0:
                for k, (lab, val) in enumerate(entries):
                    yy = 0.892 - k * 0.027
                    fig.text(xa + 0.014, yy, lab, ha="left", va="center", fontsize=10, fontweight="bold")
                    fig.text(xa + 0.115, yy, ":", ha="left", va="center", fontsize=10)
                    fig.text(xa + 0.135, yy, val, ha="left", va="center", fontsize=10)
            else:
                lab, val = entries[0]
                fig.text(xa + 0.014, 0.883, lab, ha="left", va="center", fontsize=10, fontweight="bold")
                fig.text(xa + 0.115, 0.883, ":", ha="left", va="center", fontsize=10)
                fig.text(xa + 0.135, 0.883, val, ha="left", va="center", fontsize=10)

        fig.text(0.5, 0.828,
                 "NORMAL STRESS vs PEAK SHEAR STRESS (FAILURE ENVELOPE)",
                 ha="center", va="center", fontsize=14, fontweight="bold")

        ax = fig.add_axes([0.09, 0.335, 0.87, 0.475])
        ax.set_xlim(0, xmax)
        ax.set_ylim(0, ymax)
        ax.set_xlabel("Normal Stress, σₙ (kg/cm²)", fontsize=13, fontweight="bold", labelpad=8)
        ax.xaxis.set_label_coords(0.5, -0.035)
        ax.set_ylabel("Peak Shear Stress, τ (kg/cm²)", fontsize=13, fontweight="bold")
        ax.grid(True, linestyle="--", linewidth=0.7, alpha=0.45)
        ax.tick_params(labelsize=10)
        ax.set_xticks([round(i * x_step, 10) for i in range(int(round(xmax / x_step)) + 1)])
        ax.set_yticks([round(i * y_step, 10) for i in range(int(round(ymax / y_step)) + 1)])

        first_x = min(sigma)
        first_y = calc_c + slope * first_x
        ax.plot([0, first_x], [calc_c, first_y], color="red", linewidth=2.0, linestyle="--")
        ax.plot([first_x, xmax], [first_y, calc_c + slope * xmax],
                color="red", linewidth=2.0, label="Failure Envelope")
        ax.scatter(sigma, tau, s=105, facecolors="white", edgecolors="black",
                   linewidths=2.0, zorder=5, label="Test Results (Peak)")

        for x, yy in zip(sigma, tau):
            ax.annotate(f"({x:.2f}, {yy:.2f})", (x, yy),
                        xytext=(0, 14), textcoords="offset points", ha="center",
                        fontsize=10, fontweight="bold", color="#1238a5")

        handles, labels = ax.get_legend_handles_labels()
        if len(handles) == 2:
            ax.legend([handles[1], handles[0]], ["Test Results (Peak)", "Failure Envelope"],
                      loc="upper left", fontsize=10, frameon=True,
                      fancybox=False, edgecolor="black")

        eq = ("τ = c + σₙ tan φ\n"
              f"    = {calc_c:.2f} + σₙ tan {calc_phi:.1f}°\n\n"
              f"c = {calc_c:.2f} kg/cm²\n"
              f"φ = {calc_phi:.1f}°\n\n"
              f"R² = {r2:.3f}")
        ax.text(0.985, 0.035, eq, transform=ax.transAxes,
                ha="right", va="bottom", fontsize=10.5, fontweight="bold",
                color="black", linespacing=1.45,
                bbox=dict(boxstyle="square,pad=0.65", facecolor="white",
                          edgecolor="black", linewidth=1.0))

        fig.text(0.5, 0.275, "DIRECT SHEAR TEST RESULTS",
                 ha="center", va="center", fontsize=14,
                 fontweight="bold", color="#1425a0")

        bottom_ax = fig.add_axes([0.008, 0.065, 0.984, 0.205])
        bottom_ax.axis("off")
        headers = [
            "Sr. No.", "Sample Depth\n(m)", "Sample Type", "Density\n(gm/cc)",
            "Normal Stress,\nσₙ (kg/cm²)", "Shear Stress at Failure\n(Peak), τ (kg/cm²)",
            "Cohesion,\nc (kg/cm²)", "Angle of Internal Friction,\nφ (°)"
        ]
        body = []
        for i, s_val, t_val in zip(range(1, len(sigma) + 1), sigma, tau):
            body.append([
                str(i), f"{float(result['depth']):.2f}" if i == 2 else "",
                result["sample_type"] if i == 2 else "",
                f"{float(result['density']):.2f}" if i == 2 else "",
                f"{float(s_val):.2f}", f"{float(t_val):.2f}",
                f"{calc_c:.2f}" if i == 2 else "",
                f"{calc_phi:.1f}" if i == 2 else ""
            ])

        table = bottom_ax.table(
            cellText=body, colLabels=headers, cellLoc="center", colLoc="center",
            loc="center", colWidths=[0.07, 0.13, 0.13, 0.11, 0.14, 0.17, 0.105, 0.145]
        )
        table.auto_set_font_size(False)
        table.set_fontsize(9.2)
        table.scale(1, 2.25)
        merged_cols = {1, 2, 3, 6, 7}
        for (r, c), cell in table.get_celld().items():
            cell.set_edgecolor("black")
            cell.set_linewidth(0.8)
            cell.set_facecolor("white")
            cell.PAD = 0.02
            if r == 0:
                cell.set_facecolor("#eef3fb")
                cell.get_text().set_fontweight("bold")
            elif c in merged_cols:
                if r in (1, 3):
                    cell.get_text().set_text("")
                if r == 1:
                    cell.visible_edges = "TLR"
                elif r == 2:
                    cell.visible_edges = "LR"
                elif r == 3:
                    cell.visible_edges = "BLR"

        filename = f"{safe_name(bh_id)}_Depth_{float(result['depth']):.2f}m_Failure_Envelope.png"
        out_file = plot_dir / filename
        fig.savefig(out_file, dpi=200, bbox_inches="tight", pad_inches=0.04,
                    facecolor="white", edgecolor="none")
        plt.close(fig)
        created_files.append(out_file.name)

    return created_files

@app.route("/api/dst/export_excel", methods=["POST"])
def api_dst_export_excel():
    import zipfile
    import base64

    data = request.json or {}
    results_raw = data.get("results", [])
    bh_id = data.get("bh_id", "DST_Result").strip() or "DST_Result"
    pr = float(data.get("pr_constant", 0.325))
    stresses = [float(data.get("normal_1", 0.5)), float(data.get("normal_2", 1.0)), float(data.get("normal_3", 1.5))]

    template_path = DST_DIR / "DST Template.xlsx"
    if not template_path.exists():
        return jsonify({"error": "DST Template.xlsx not found."}), 404

    reconstructed_results = []
    for r in results_raw:
        trials_objs = []
        for t in r["trials"]:
            trial_obj = dst_generator.DSTTrial(
                normal_stress=t["normal_stress"], displacement=t["displacement"],
                shear_stress=t["shear_stress"], failure_shear_stress=t["failure_shear_stress"],
                peak_shear_stress=t["peak_shear_stress"], peak_displacement=t["peak_displacement"],
                residual_shear_stress=t["residual_shear_stress"]
            )
            trials_objs.append(trial_obj)
        
        reconstructed_results.append({
            "depth": float(r["depth"]), "soil": r["soil"], "sample_type": r["sample_type"],
            "density": float(r["density"]), "weight": float(r["weight"]), "c": float(r["c"]),
            "phi": float(r["phi"]), "trials": trials_objs, "summary": r["summary"]
        })

    outputs_base = BASE_DIR / "outputs" / "DST"
    outputs_base.mkdir(parents=True, exist_ok=True)

    safe_bh = "".join(c if c.isalnum() or c in "-_" else "_" for c in bh_id) or "DST_Result"
    excel_filename = f"{safe_bh}_DST_MultiDepth.xlsx"
    excel_path = outputs_base / excel_filename
    plot_dir = outputs_base / f"{safe_bh}_Failure_Plots"

    try:
        # 1. Generate Excel file
        excel_export_multi.export_multiple_dst_to_excel(
            template_path=template_path, output_path=str(excel_path), bh_id=bh_id,
            pr_constant=pr, normal_stresses=stresses, results=reconstructed_results
        )

        # 2. Generate all Failure Envelope Plots (PNGs)
        _render_dst_plots(bh_id, reconstructed_results, plot_dir)

        # 3. Create ZIP Package matching LL & PL structure
        zip_filename = f"DST_Package_{safe_bh}.zip"
        zip_path = outputs_base / zip_filename
        with zipfile.ZipFile(zip_path, "w", zipfile.ZIP_DEFLATED) as zip_file:
            if excel_path.exists():
                zip_file.write(excel_path, arcname=excel_path.name)
            if plot_dir.exists():
                for item in plot_dir.iterdir():
                    if item.is_file():
                        zip_file.write(item, arcname=f"{plot_dir.name}/{item.name}")

        # 4. Read ZIP base64 for instant client browser download
        with open(zip_path, "rb") as f:
            b64_zip = base64.b64encode(f.read()).decode("utf-8")

        return jsonify({
            "success": True,
            "count": len(reconstructed_results),
            "excel_path": str(excel_path),
            "excel_filename": excel_filename,
            "plot_dir": str(plot_dir),
            "zip_filename": zip_filename,
            "zip_path": str(zip_path),
            "download_zip_url": f"/outputs/DST/{zip_filename}",
            "download_excel_url": f"/outputs/DST/{excel_filename}",
            "file_base64": b64_zip
        })
    except Exception as e:
        return jsonify({"error": str(e)}), 500

@app.route("/api/dst/export_plots", methods=["POST"])
def api_dst_export_plots():
    import zipfile
    import base64

    data = request.json or {}
    results_raw = data.get("results", [])
    bh_id = data.get("bh_id", "DST_Result").strip() or "DST_Result"

    if not results_raw:
        return jsonify({"error": "No DST results to plot."}), 400

    outputs_base = BASE_DIR / "outputs" / "DST"
    outputs_base.mkdir(parents=True, exist_ok=True)

    safe_bh = "".join(c if c.isalnum() or c in "-_" else "_" for c in bh_id) or "DST_Result"
    plot_dir = outputs_base / f"{safe_bh}_Failure_Plots"

    try:
        created_files = _render_dst_plots(bh_id, results_raw, plot_dir)

        # Create ZIP package of all plots
        zip_filename = f"{safe_bh}_Failure_Plots.zip"
        zip_path = outputs_base / zip_filename
        with zipfile.ZipFile(zip_path, "w", zipfile.ZIP_DEFLATED) as zipf:
            for item in plot_dir.iterdir():
                if item.is_file() and item.name.endswith(".png"):
                    zipf.write(item, arcname=f"{plot_dir.name}/{item.name}")

        with open(zip_path, "rb") as f:
            b64_zip = base64.b64encode(f.read()).decode("utf-8")

        return jsonify({
            "success": True,
            "count": len(created_files),
            "output_dir": str(plot_dir),
            "zip_filename": zip_filename,
            "download_zip_url": f"/outputs/DST/{zip_filename}",
            "file_base64": b64_zip,
            "files": created_files
        })
    except Exception as e:
        return jsonify({"error": str(e)}), 500

@app.route("/api/dst/import_excel", methods=["POST"])
def api_dst_import_excel():
    if "file" not in request.files: return jsonify({"error": "No file uploaded"}), 400
    file = request.files["file"]
    try:
        wb = load_workbook(file, data_only=True)
        if "Input" not in wb.sheetnames: return jsonify({"error": "Worksheet 'Input' was not found in the selected Excel file."}), 400
        ws = wb["Input"]
        bh = ws["B2"].value
        pr = ws["B3"].value
        n1 = ws["B4"].value
        n2 = ws["B5"].value
        n3 = ws["B6"].value
        seed = ws["B7"].value

        if bh is None or str(bh).strip() == "":
            return jsonify({"error": "BH ID is required in cell B2."}), 400

        headers = {}
        for col in range(1, ws.max_column + 1):
            val = ws.cell(10, col).value
            if val is not None:
                headers[str(val).strip().lower()] = col

        required = ["depth (m)", "soil type", "sample type", "density (g/cc)", "target c (kg/cm²)", "target φ (°)"]
        missing = [h for h in required if h not in headers]
        if missing:
            return jsonify({"error": f"Missing required Excel column(s): {', '.join(missing)}"}), 400

        rows = []
        for r in range(11, ws.max_row + 1):
            vals = {k: ws.cell(r, col).value for k, col in headers.items()}
            if all(v is None or str(v).strip() == "" for v in vals.values()):
                continue

            depth = vals.get("depth (m)")
            soil = vals.get("soil type")
            sample_type = vals.get("sample type")
            density = vals.get("density (g/cc)")
            c = vals.get("target c (kg/cm²)")
            phi = vals.get("target φ (°)")

            if any(v is None or str(v).strip() == "" for v in (depth, soil, sample_type, density, c, phi)):
                return jsonify({"error": f"Excel row {r}: Depth, Soil Type, Sample Type, Density, Target c and Target φ are required."}), 400

            d_val = float(density)
            rows.append({
                "depth": float(depth),
                "soil": str(soil).strip(),
                "sample_type": str(sample_type).strip(),
                "density": d_val,
                "weight": round(d_val * 90.0, 2),
                "c": float(c),
                "phi": float(phi),
                "behaviour": dst_generator.classify_density(str(soil).strip(), d_val)
            })

        if not rows:
            return jsonify({"error": "No depth data was found. Enter data from row 11 onward."}), 400

        return jsonify({
            "success": True,
            "bh_id": str(bh).strip(),
            "pr_constant": str(pr or "0.325"),
            "normal_1": str(n1 or "0.50"),
            "normal_2": str(n2 or "1.00"),
            "normal_3": str(n3 or "1.50"),
            "random_seed": "AUTO" if seed is None or str(seed).strip() == "" else str(seed),
            "rows": rows
        })
    except Exception as e:
        return jsonify({"error": str(e)}), 400

# =========================================================================
# FSI (Free Swell Index) Backend Endpoints - Exact Existing Logic
# =========================================================================

@app.route("/api/fsi/options", methods=["GET"])
def api_fsi_options():
    return jsonify({"soil_options": fsi_engine.SOIL_OPTIONS, "sample_types": fsi_engine.SAMPLE_TYPES})

@app.route("/api/fsi/generate", methods=["POST"])
def api_fsi_generate():
    data = request.json or {}
    bh = data.get("bh_id", "").strip() or "RBH-12"
    depth_rows = data.get("rows", [])
    if not depth_rows:
        return jsonify({"error": "Please enter at least one depth."}), 400

    used_depths = set()
    for i, item in enumerate(depth_rows, start=1):
        try:
            d_val = float(item.get("depth", 0))
            if d_val < 0:
                return jsonify({"error": f"Sample Depth in row {i} cannot be negative."}), 400
            if d_val in used_depths:
                return jsonify({"error": f"Duplicate Sample Depth {d_val:g} m."}), 400
            used_depths.add(d_val)
        except (ValueError, TypeError):
            return jsonify({"error": f"Invalid Sample Depth in row {i}."}), 400

    results = []
    for i, item in enumerate(depth_rows, start=1):
        depth_val = float(item.get("depth", 1.5))
        soil_choice = item.get("soil", "Auto / Random Soil Type")
        soil = fsi_engine.choose_soil(soil_choice)
        sample_type = item.get("sample_type", "UDS")

        vk, vd, fsi = fsi_engine.generate_fsi_readings(soil)
        results.append({
            "SR No": i,
            "BH ID": bh,
            "Sample ID": f"FSI-{i:02d}",
            "Sample Type": sample_type,
            "Depth": round(depth_val, 2),
            "Soil Description": soil,
            "Vk": int(vk),
            "Vd": int(vd),
            "Difference": int(vd - vk),
            "FSI": round(((vd - vk) / vk) * 100.0, 2)
        })
    return jsonify({"success": True, "results": results})

@app.route("/api/fsi/export_excel", methods=["POST"])
def api_fsi_export_excel():
    data = request.json or {}
    rows = data.get("results", [])
    bh = data.get("bh_id", "RBH-12").strip() or "RBH-12"

    # If results is empty but input rows are provided, auto-generate
    if not rows and data.get("rows"):
        depth_rows = data.get("rows", [])
        results = []
        for i, item in enumerate(depth_rows, start=1):
            depth_val = float(item.get("depth", 1.5))
            soil_choice = item.get("soil", "Auto / Random Soil Type")
            soil = fsi_engine.choose_soil(soil_choice)
            sample_type = item.get("sample_type", "UDS")
            vk, vd, fsi = fsi_engine.generate_fsi_readings(soil)
            results.append({
                "SR No": i,
                "BH ID": bh,
                "Sample ID": f"FSI-{i:02d}",
                "Sample Type": sample_type,
                "Depth": round(depth_val, 2),
                "Soil Description": soil,
                "Vk": int(vk),
                "Vd": int(vd),
                "Difference": int(vd - vk),
                "FSI": round(((vd - vk) / vk) * 100.0, 2)
            })
        rows = results

    if not rows:
        return jsonify({"error": "No data to export. Please generate FSI data first."}), 400

    clean_rows = []
    for r in rows:
        clean_rows.append({
            "SR No": int(r.get("SR No", 1)),
            "BH ID": str(r.get("BH ID", bh)),
            "Sample ID": str(r.get("Sample ID", f"FSI-{int(r.get('SR No', 1)):02d}")),
            "Sample Type": str(r.get("Sample Type", "UDS")),
            "Depth": float(r.get("Depth", 1.5)),
            "Soil Description": str(r.get("Soil Description", "")),
            "Vk": int(r.get("Vk", 10)),
            "Vd": int(r.get("Vd", 15)),
            "Difference": int(r.get("Difference", int(r.get("Vd", 15)) - int(r.get("Vk", 10)))),
            "FSI": float(r.get("FSI", round(((int(r.get("Vd", 15)) - int(r.get("Vk", 10))) / int(r.get("Vk", 10))) * 100.0, 2)))
        })

    outputs_base = BASE_DIR / "outputs" / "FSI"
    outputs_base.mkdir(parents=True, exist_ok=True)
    safe_bh = "".join(c if c.isalnum() or c in "-_" else "_" for c in bh) or "FSI_Result"
    excel_filename = f"FSI_Report_{safe_bh}.xlsx"
    excel_path = outputs_base / excel_filename

    try:
        fsi_engine.export_excel(clean_rows, str(excel_path))
        import zipfile
        import base64

        # Create ZIP package matching DST and LL&PL
        zip_filename = f"FSI_Package_{safe_bh}.zip"
        zip_path = outputs_base / zip_filename
        with zipfile.ZipFile(zip_path, "w", zipfile.ZIP_DEFLATED) as zip_file:
            if excel_path.exists():
                zip_file.write(excel_path, arcname=excel_path.name)

        with open(excel_path, "rb") as f:
            b64_excel = base64.b64encode(f.read()).decode("utf-8")

        with open(zip_path, "rb") as f:
            b64_zip = base64.b64encode(f.read()).decode("utf-8")

        return jsonify({
            "success": True,
            "count": len(clean_rows),
            "excel_path": str(excel_path),
            "excel_filename": excel_filename,
            "zip_path": str(zip_path),
            "zip_filename": zip_filename,
            "download_url": f"/outputs/FSI/{excel_filename}",
            "download_zip_url": f"/outputs/FSI/{zip_filename}",
            "excel_base64": b64_excel,
            "zip_base64": b64_zip,
            "file_base64": b64_zip
        })
    except Exception as e:
        return jsonify({"error": str(e)}), 500

@app.route("/api/fsi/download_excel_direct", methods=["POST", "GET"])
def api_fsi_download_excel_direct():
    import json
    if request.method == "POST":
        data_str = request.form.get("data")
        if data_str:
            try:
                payload = json.loads(data_str)
            except Exception:
                payload = {}
        else:
            payload = request.json or {}
    else:
        payload = request.args

    rows = payload.get("results", [])
    bh = payload.get("bh_id", "RBH-12")
    if isinstance(bh, str):
        bh = bh.strip() or "RBH-12"

    if not rows and payload.get("rows"):
        depth_rows = payload.get("rows", [])
        results = []
        for i, item in enumerate(depth_rows, start=1):
            depth_val = float(item.get("depth", 1.5))
            soil_choice = item.get("soil", "Auto / Random Soil Type")
            soil = fsi_engine.choose_soil(soil_choice)
            sample_type = item.get("sample_type", "UDS")
            vk, vd, fsi = fsi_engine.generate_fsi_readings(soil)
            results.append({
                "SR No": i,
                "BH ID": bh,
                "Sample ID": f"FSI-{i:02d}",
                "Sample Type": sample_type,
                "Depth": round(depth_val, 2),
                "Soil Description": soil,
                "Vk": int(vk),
                "Vd": int(vd),
                "Difference": int(vd - vk),
                "FSI": round(((vd - vk) / vk) * 100.0, 2)
            })
        rows = results

    if not rows:
        return "No data to export.", 400

    clean_rows = []
    for r in rows:
        clean_rows.append({
            "SR No": int(r.get("SR No", 1)),
            "BH ID": str(r.get("BH ID", bh)),
            "Sample ID": str(r.get("Sample ID", f"FSI-{int(r.get('SR No', 1)):02d}")),
            "Sample Type": str(r.get("Sample Type", "UDS")),
            "Depth": float(r.get("Depth", 1.5)),
            "Soil Description": str(r.get("Soil Description", "")),
            "Vk": int(r.get("Vk", 10)),
            "Vd": int(r.get("Vd", 15)),
            "Difference": int(r.get("Difference", int(r.get("Vd", 15)) - int(r.get("Vk", 10)))),
            "FSI": float(r.get("FSI", round(((int(r.get("Vd", 15)) - int(r.get("Vk", 10))) / int(r.get("Vk", 10))) * 100.0, 2)))
        })

    outputs_base = BASE_DIR / "outputs" / "FSI"
    outputs_base.mkdir(parents=True, exist_ok=True)
    safe_bh = "".join(c if c.isalnum() or c in "-_" else "_" for c in str(bh)) or "FSI_Result"
    excel_filename = f"FSI_Report_{safe_bh}.xlsx"
    excel_path = outputs_base / excel_filename

    try:
        fsi_engine.export_excel(clean_rows, str(excel_path))
        return send_file(
            str(excel_path),
            as_attachment=True,
            download_name=excel_filename,
            mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
        )
    except Exception as e:
        return str(e), 500


# =========================================================================
# LL & PL (Liquid Limit & Plastic Limit) - Exact Master Script Engine
# =========================================================================

@app.route("/api/ll_pl/generate", methods=["POST"])
def api_ll_pl_generate():
    data = request.json or {}
    rows = data.get("rows", [])
    bh = data.get("bh_id", "").strip() or "BH-01"
    chainage = data.get("chainage", "").strip()

    valid_samples = []
    for i, r in enumerate(rows, start=1):
        depth_str = str(r.get("depth", "")).strip()
        ll_str = str(r.get("ll", "")).strip()
        pl_str = str(r.get("pl", "")).strip()
        sample_type = str(r.get("sample_type", "DS")).strip() or "DS"

        if not depth_str and not ll_str and not pl_str:
            continue
        if not depth_str:
            return jsonify({"error": f"Row {i}: Depth is required."}), 400
        if not ll_str:
            return jsonify({"error": f"Row {i}: Liquid Limit (LL) is required."}), 400

        try:
            depth_val = float(depth_str)
            ll_val = float(ll_str)
            if ll_val <= 0 or ll_val >= 100:
                return jsonify({"error": f"Row {i}: Liquid Limit (LL) must be between 0 and 100."}), 400
        except ValueError:
            return jsonify({"error": f"Row {i}: Values must be numbers."}), 400

        pl_val = None
        if pl_str:
            try:
                pl_val = float(pl_str)
                if pl_val <= 0 or pl_val >= 100:
                    return jsonify({"error": f"Row {i}: Plastic Limit (PL) must be between 0 and 100."}), 400
                if pl_val >= ll_val:
                    return jsonify({"error": f"Row {i}: Plastic Limit (PL) must be less than Liquid Limit (LL)."}), 400
            except ValueError:
                return jsonify({"error": f"Row {i}: Plastic Limit must be a number."}), 400

        valid_samples.append({
            "sr": len(valid_samples) + 1, "bh": bh, "chainage": chainage,
            "depth": round(depth_val, 2), "sample_type": sample_type,
            "ll": round(ll_val, 2), "pl": round(pl_val, 2) if pl_val is not None else None
        })

    if not valid_samples:
        return jsonify({"error": "Please enter at least one sample."}), 400

    results = []
    for sample in valid_samples:
        blows, moisture, calc_ll, r2 = llpl_engine.generate_ll_observations(sample["ll"])
        slope, intercept, _, _ = llpl_engine.regression(blows, moisture)

        pl_calc = sample["pl"]
        pi_calc = round(calc_ll - pl_calc, 2) if pl_calc is not None else None

        a_line = 0.73 * (calc_ll - 20.0)
        if pi_calc is not None:
            if calc_ll < 35:
                soil_class = "CL (Low Plasticity Clay)" if pi_calc >= a_line else "ML (Low Plasticity Silt)"
            elif calc_ll <= 50:
                soil_class = "CI (Intermediate Plasticity Clay)" if pi_calc >= a_line else "MI (Intermediate Plasticity Silt)"
            else:
                soil_class = "CH (High Plasticity Clay)" if pi_calc >= a_line else "MH (High Plasticity Silt)"
        else:
            soil_class = "Non-Plastic / Liquid Limit only"

        results.append({
            "sr": sample["sr"], "bh": sample["bh"], "chainage": sample["chainage"],
            "depth": sample["depth"], "sample_type": sample["sample_type"],
            "input_ll": sample["ll"], "input_pl": sample["pl"],
            "blows": blows, "moisture": [round(m, 2) for m in moisture],
            "slope": slope, "intercept": intercept,
            "calculated_ll": round(calc_ll, 2), "calculated_pl": round(pl_calc, 2) if pl_calc is not None else None,
            "calculated_pi": pi_calc, "r2": round(r2, 3), "soil_class": soil_class
        })

    return jsonify({"success": True, "results": results})

@app.route("/api/ll_pl/export_excel", methods=["POST"])
def api_ll_pl_export_excel():
    data = request.json or {}
    rows = data.get("rows", [])
    bh = data.get("bh_id", "").strip()
    chainage = data.get("chainage", "").strip()

    if not bh:
        return jsonify({"error": "Borehole No. is required."}), 400
    if not chainage:
        return jsonify({"error": "Chainage is required."}), 400

    valid_samples = []
    for i, r in enumerate(rows, start=1):
        depth_str = str(r.get("depth", "")).strip()
        ll_str = str(r.get("ll", "")).strip()
        pl_str = str(r.get("pl", "")).strip() if r.get("pl") is not None else ""
        sample_type = str(r.get("sample_type", "DS")).strip() or "DS"

        if not depth_str and not ll_str and not pl_str:
            continue
        if not depth_str:
            return jsonify({"error": f"Row {i}: Depth is required."}), 400
        if not ll_str:
            return jsonify({"error": f"Row {i}: Liquid Limit (LL) is required."}), 400

        try:
            depth_val = float(depth_str)
        except ValueError:
            return jsonify({"error": f"Row {i}: Depth must be a number."}), 400

        if depth_val < 0:
            return jsonify({"error": f"Row {i}: Depth cannot be negative."}), 400

        try:
            ll_val = float(ll_str)
        except ValueError:
            return jsonify({"error": f"Row {i}: Liquid Limit (LL) must be a number."}), 400

        if ll_val <= 0 or ll_val >= 100:
            return jsonify({"error": f"Row {i}: Liquid Limit (LL) must be between 0 and 100."}), 400

        pl_val = None
        if pl_str:
            try:
                pl_val = float(pl_str)
            except ValueError:
                return jsonify({"error": f"Row {i}: Plastic Limit (PL) must be a number."}), 400

            if pl_val <= 0 or pl_val >= 100:
                return jsonify({"error": f"Row {i}: Plastic Limit (PL) must be between 0 and 100."}), 400
            if pl_val >= ll_val:
                return jsonify({"error": f"Row {i}: Plastic Limit (PL) must be less than Liquid Limit (LL)."}), 400

        valid_samples.append({
            "sr": len(valid_samples) + 1, "bh": bh, "chainage": chainage,
            "depth": round(depth_val, 2), "sample_type": sample_type,
            "ll": round(ll_val, 2), "pl": round(pl_val, 2) if pl_val is not None else None
        })

    if not valid_samples:
        return jsonify({"error": "Please enter at least one sample."}), 400

    import zipfile
    import shutil

    # Ensure output directory exists in workspace
    outputs_base = BASE_DIR / "outputs" / "LL_PL"
    outputs_base.mkdir(parents=True, exist_ok=True)

    safe_bh = "".join(c if c.isalnum() or c in "-_" else "_" for c in bh) or "Generated"
    excel_filename = f"LL_PL_Generated_{safe_bh}.xlsx"
    excel_path = outputs_base / excel_filename

    try:
        # Calls the exact master function from original script
        # This creates:
        # 1. excel_path (.xlsx)
        # 2. excel_path.parent / f"{excel_path.stem}_Flow_Plots" (folder with all PNG flow curves)
        # 3. excel_path.parent / f"{excel_path.stem}_Flow_Plots" / f"{excel_path.stem}_Flow_Plots.pdf"
        results = llpl_engine.generate_from_samples(None, str(excel_path), valid_samples)
        
        # Ensure 'Input Sheet' is removed from the saved workbook
        if excel_path.exists():
            try:
                from openpyxl import load_workbook
                wb_chk = load_workbook(str(excel_path))
                if "Input Sheet" in wb_chk.sheetnames and len(wb_chk.sheetnames) > 1:
                    del wb_chk["Input Sheet"]
                    wb_chk.save(str(excel_path))
            except Exception:
                pass
        
        plot_dir = Path(results[0].get("plot_dir", str(outputs_base / f"{excel_path.stem}_Flow_Plots"))) if results else outputs_base / f"{excel_path.stem}_Flow_Plots"
        plot_pdf = Path(results[0].get("plot_pdf", str(plot_dir / f"{excel_path.stem}_Flow_Plots.pdf"))) if results else plot_dir / f"{excel_path.stem}_Flow_Plots.pdf"

        # Create a ZIP package containing the Excel file, the entire Flow Plots folder (PNGs), and PDF
        zip_path = outputs_base / f"LL_PL_Package_{safe_bh}.zip"
        with zipfile.ZipFile(zip_path, "w", zipfile.ZIP_DEFLATED) as zip_file:
            if excel_path.exists():
                zip_file.write(excel_path, arcname=excel_path.name)
            if plot_dir.exists():
                for item in plot_dir.iterdir():
                    if item.is_file():
                        zip_file.write(item, arcname=f"{plot_dir.name}/{item.name}")

        import base64
        file_base64 = None
        if zip_path.exists():
            try:
                with open(zip_path, "rb") as zf:
                    file_base64 = base64.b64encode(zf.read()).decode("utf-8")
            except Exception:
                file_base64 = None

        return jsonify({
            "success": True,
            "count": len(results),
            "excel_path": str(excel_path),
            "excel_filename": excel_filename,
            "plot_dir": str(plot_dir),
            "plot_pdf": str(plot_pdf),
            "zip_filename": zip_path.name,
            "file_base64": file_base64,
            "download_excel_url": f"/outputs/LL_PL/{excel_filename}",
            "download_zip_url": f"/outputs/LL_PL/{zip_path.name}",
            "download_pdf_url": f"/outputs/LL_PL/{plot_dir.name}/{plot_pdf.name}" if plot_pdf.exists() else None
        })
    except Exception as error:
        return jsonify({"error": str(error)}), 500

# =========================================================================
# HYD (Hydrometer) Backend Endpoints - Exact PSD & Synthetic Model
# =========================================================================

import hydrometer_synthetic_generator_v1_9_auto_clay_PSD_v2_9_plot_GSD_Fines_fixed as hyd_engine

@app.route("/api/hyd/options", methods=["GET"])
def api_hyd_options():
    return jsonify({
        "presets": list(hyd_engine.PRESETS.keys()),
        "obs_times": hyd_engine.OBS_TIMES,
        "calibration": hyd_engine.CALIBRATION
    })

@app.route("/api/hyd/parse_template", methods=["POST"])
def api_hyd_parse_template():
    if "file" not in request.files:
        return jsonify({"error": "No template file uploaded."}), 400
    file = request.files["file"]
    template_type = request.form.get("template_type", "Type 1")

    try:
        cache_dir = BASE_DIR / "outputs" / "HYD" / "cache"
        cache_dir.mkdir(parents=True, exist_ok=True)
        filename = file.filename or "GSD_Template.xlsx"
        cached_path = cache_dir / filename
        file.save(str(cached_path))

        keep_vba = filename.lower().endswith(".xlsm")
        wb = load_workbook(str(cached_path), data_only=True)
        if "GSD" not in wb.sheetnames or "HYD" not in wb.sheetnames:
            return jsonify({"error": "GSD or HYD sheet not found in the uploaded workbook."}), 400

        ws_gsd = wb["GSD"]
        bh_id = str(ws_gsd["B5"].value or "").strip() or "BH-01"

        if template_type == "Type 1":
            gravel_col, sand_col, fines_col = 27, 28, 29  # AA, AB, AC
        else:
            gravel_col, sand_col, fines_col = 23, 24, 25  # W, X, Y

        rows = []
        for r in range(5, ws_gsd.max_row + 1):
            depth = ws_gsd.cell(r, 6).value
            sample_type = ws_gsd.cell(r, 5).value
            gravel = ws_gsd.cell(r, gravel_col).value
            sand = ws_gsd.cell(r, sand_col).value
            fines = ws_gsd.cell(r, fines_col).value

            if all(v is None or str(v).strip() == "" for v in (depth, sample_type, gravel, sand, fines)):
                break

            def _clean_str(v):
                if v is None or str(v).strip() == "":
                    return ""
                try:
                    return f"{float(v):.2f}"
                except Exception:
                    return str(v)

            rows.append({
                "gsd_row": r,
                "depth": str(depth if depth is not None else f"{1.50 + (len(rows))*1.50:.2f}"),
                "sample_type": str(sample_type or "UDS"),
                "gravel": _clean_str(gravel),
                "sand": _clean_str(sand),
                "fines": _clean_str(fines),
                "silt": "",
                "clay": "",
                "preset": "2 - Normal Silt",
                "generate": False,
                "spgr": "",
                "hydrometer": "",
                "cylinder": ""
            })

        wb.close()

        return jsonify({
            "success": True,
            "bh_id": bh_id,
            "template_type": template_type,
            "cached_file": filename,
            "rows": rows
        })
    except Exception as e:
        return jsonify({"error": str(e)}), 400

@app.route("/api/hyd/generate", methods=["POST"])
def api_hyd_generate():
    import json
    import zipfile
    import time

    if request.is_json:
        data = request.json or {}
    else:
        config_str = request.form.get("config", "{}")
        try:
            data = json.loads(config_str)
        except Exception:
            data = {}

    rows = data.get("rows", [])
    bh_id = data.get("bh_id", "BH-01").strip() or "BH-01"
    template_type = data.get("template_type", "Type 1")
    wb_weight = float(data.get("wb", data.get("wb_weight", 50.0)))
    tolerance = float(data.get("tolerance", 0.75))
    seed_str = str(data.get("seed", data.get("random_seed", 20260819))).strip()
    seed_val = int(float(seed_str)) if seed_str and seed_str.replace('.', '', 1).isdigit() else 20260819
    rng = random.Random(seed_val)

    if not rows:
        return jsonify({"error": "No sample rows loaded."}), 400

    outputs_base = BASE_DIR / "outputs" / "HYD"
    outputs_base.mkdir(parents=True, exist_ok=True)
    cache_dir = outputs_base / "cache"
    cache_dir.mkdir(parents=True, exist_ok=True)

    # Check if a file was uploaded in this request
    uploaded_file = request.files.get("file")
    if uploaded_file and uploaded_file.filename:
        cached_file = uploaded_file.filename
        cache_path = cache_dir / cached_file
        uploaded_file.save(str(cache_path))
    else:
        cached_file = data.get("cached_file", "")
        cache_path = cache_dir / cached_file if cached_file else None

    # Timestamped subfolder for clean storage and isolation
    run_timestamp = time.strftime("%Y%m%d_%H%M%S")
    run_dir = outputs_base / f"run_{run_timestamp}_{bh_id}"
    run_dir.mkdir(parents=True, exist_ok=True)
    plot_dir = run_dir / "PSD_Plots"
    plot_dir.mkdir(parents=True, exist_ok=True)

    if cache_path and cache_path.exists():
        keep_vba = str(cache_path).lower().endswith(".xlsm")
        wb = load_workbook(str(cache_path), keep_vba=keep_vba)
        wb_data = load_workbook(str(cache_path), data_only=True)
        ws_gsd_values = wb_data["GSD"]
        base_name, ext = os.path.splitext(cache_path.name)
        excel_filename = f"{base_name}_HYD_Generated{ext}"
        excel_path = run_dir / excel_filename
        hyd_sheet = wb["HYD"] if "HYD" in wb.sheetnames else None
    else:
        from openpyxl import Workbook
        wb = Workbook()
        ws_gsd_values = wb.active
        ws_gsd_values.title = "GSD"
        excel_filename = f"{bh_id}_HYD_Generated.xlsx"
        excel_path = run_dir / excel_filename
        hyd_sheet = wb.create_sheet("HYD")
        hyd_sheet.append(["SR", "BH ID", "Sample ID", "Depth", "Temp", "Viscosity", "SPGR", "Cm", "Hydro No", "Cylinder", "Mt", "X"] + [f"R_{t}" for t in hyd_engine.OBS_TIMES])
        wb_data = wb

    results = []
    summary_rows = []
    plots_list = []
    hyd_generated_count = 0

    for idx, row in enumerate(rows, 1):
        fines_str = str(row.get("fines", "100.0")).strip()
        fines = float(fines_str) if fines_str else 100.0
        depth = str(row.get("depth", f"{idx * 1.50:.2f}")).strip() or f"{idx * 1.50:.2f}"
        sample_type = str(row.get("sample_type", "UDS")).strip() or "UDS"
        gsd_row = int(row.get("gsd_row", idx + 4))
        should_gen_hyd = bool(row.get("generate", False))

        sample_res = None
        if should_gen_hyd:
            hyd_generated_count += 1

            silt_val = str(row.get("silt", "")).strip()
            silt = float(silt_val) if silt_val else fines * 0.6
            clay = max(0.0, fines - silt)
            preset = str(row.get("preset", "2 - Normal Silt")).strip() or "2 - Normal Silt"

            spgr_val = str(row.get("spgr", "")).strip()
            spgr = float(spgr_val) if spgr_val else 2.65

            hydro_val = str(row.get("hydrometer", "")).strip()
            hydro = int(float(hydro_val)) if hydro_val else 1

            cyl_val = str(row.get("cylinder", "")).strip()
            cylinder = int(float(cyl_val)) if cyl_val else 1

            class GUIVal:
                def __init__(self, val):
                    self._val = str(val) if val is not None else ""
                def get(self):
                    return self._val
                def __str__(self):
                    return self._val
                def __repr__(self):
                    return self._val
                def __float__(self):
                    return float(self._val)

            row_adapter = {
                "gsd_row": gsd_row,
                "fines": fines,
                "silt": GUIVal(silt),
                "preset": GUIVal(preset),
                "spgr": GUIVal(spgr),
                "hydrometer": GUIVal(hydro),
                "cylinder": GUIVal(cylinder),
            }

            try:
                sample_res = hyd_engine.generate_sample(row_adapter, wb_weight, tolerance, rng)

                # Update workbook HYD sheet
                if hyd_sheet is not None:
                    if cache_path and cache_path.exists():
                        hr = gsd_row
                        hyd_sheet.cell(hr, 5).value = 27.0
                        hyd_sheet.cell(hr, 6).value = sample_res["mu"]
                        hyd_sheet.cell(hr, 7).value = sample_res["gs"]
                        hyd_sheet.cell(hr, 8).value = 0.0005
                        hyd_sheet.cell(hr, 9).value = sample_res["hydrometer"]
                        hyd_sheet.cell(hr, 10).value = sample_res["cylinder"]
                        hyd_sheet.cell(hr, 11).value = 0.0
                        hyd_sheet.cell(hr, 12).value = 0.3
                        for col_idx, reading in enumerate(sample_res["readings"], start=13):
                            hyd_sheet.cell(hr, col_idx).value = reading
                            hyd_sheet.cell(hr, col_idx).number_format = "0.000"
                    else:
                        hyd_sheet.append([
                            idx, bh_id, f"HYD-{idx:02d}", depth, 27.0, sample_res["mu"], sample_res["gs"],
                            0.0005, sample_res["hydrometer"], sample_res["cylinder"], 0.0, 0.3
                        ] + sample_res["readings"])

                # Points data for frontend preview
                points = []
                for t, reading, d, p, ideal in sample_res["points"]:
                    points.append({
                        "time": t, "reading": reading, "diameter": round(d, 5) if d else None,
                        "percent_finer": round(p, 2), "ideal_reading": round(ideal, 4)
                    })

                results.append({
                    "row_index": idx,
                    "gsd_row": gsd_row,
                    "depth": depth,
                    "sample_type": sample_type,
                    "fines": fines,
                    "silt": silt,
                    "clay": clay,
                    "preset": preset,
                    "spgr": spgr,
                    "hydrometer": hydro,
                    "cylinder": cylinder,
                    "readings": sample_res["readings"],
                    "points": points
                })
            except Exception as e:
                return jsonify({"error": f"Row {idx} ({depth} m): {str(e)}"}), 400

        # Generate individual PSD Plot for all samples (GSD only if sample_res is None, GSD+HYD if sample_res is present)
        try:
            plot_path, plot_percentages = hyd_engine.generate_psd_plot(
                ws_gsd_values, gsd_row, sample_type, depth,
                bh_id, template_type, sample_res,
                plot_dir
            )
            plot_url_rel = f"/outputs/HYD/run_{run_timestamp}_{bh_id}/PSD_Plots/{plot_path.name}"
            plots_list.append({
                "depth": depth,
                "sample_type": sample_type,
                "filename": plot_path.name,
                "url": plot_url_rel
            })
            summary_rows.append({
                "bh_id": bh_id,
                "depth": depth,
                "sample_type": sample_type,
                "sieve_set": template_type,
                "gravel": plot_percentages[0],
                "sand": plot_percentages[1],
                "silt": plot_percentages[2],
                "clay": plot_percentages[3]
            })
        except Exception as e:
            return jsonify({"error": f"Plot generation error for depth {depth} m: {str(e)}"}), 400

    wb.save(str(excel_path))
    wb.close()
    if wb_data is not wb:
        wb_data.close()

    summary_path = hyd_engine.export_psd_summary(summary_rows, str(plot_dir))

    # Create complete all-in-one ZIP archive (Excel + PSD_Summary.csv + PSD_Plots)
    zip_filename = f"{bh_id}_HYD_Complete_Package_{run_timestamp}.zip"
    zip_path = run_dir / zip_filename
    with zipfile.ZipFile(zip_path, "w", zipfile.ZIP_DEFLATED) as zip_file:
        if excel_path.exists():
            zip_file.write(excel_path, arcname=excel_filename)
        if summary_path.exists():
            zip_file.write(summary_path, arcname="PSD_Summary.csv")
        for plot_item in plots_list:
            p_file = plot_dir / plot_item["filename"]
            if p_file.exists():
                zip_file.write(p_file, arcname=f"PSD_Plots/{plot_item['filename']}")

    run_rel = f"run_{run_timestamp}_{bh_id}"

    return jsonify({
        "success": True,
        "bh_id": bh_id,
        "total_samples": len(rows),
        "hyd_generated_samples": hyd_generated_count,
        "count": len(results),
        "excel_path": str(excel_path),
        "excel_filename": excel_filename,
        "excel_url": f"/outputs/HYD/{run_rel}/{excel_filename}",
        "zip_filename": zip_filename,
        "zip_url": f"/outputs/HYD/{run_rel}/{zip_filename}",
        "summary_filename": "PSD_Summary.csv",
        "summary_url": f"/outputs/HYD/{run_rel}/PSD_Plots/PSD_Summary.csv",
        "summary_path": str(summary_path),
        "plot_dir": str(plot_dir),
        "plots": plots_list,
        "summary_rows": summary_rows,
        "results": results
    })

# =========================================================================
# Parameter (Engineering Strength Parameter Generator) - v11 All Sheets Engine
# =========================================================================

import engineering_strength_parameter_generator_v11_formula_safe_all_sheets as param_engine

@app.route("/api/parameter/inspect_file", methods=["POST"])
def api_parameter_inspect_file():
    if "file" not in request.files:
        return jsonify({"error": "No file uploaded."}), 400
    file = request.files["file"]
    if not file.filename:
        return jsonify({"error": "Empty filename."}), 400

    try:
        wb = load_workbook(file, read_only=True)
        sheets = wb.sheetnames
        wb.close()

        base = Path(file.filename).stem
        default_output = f"{base}_Synthetic_Strength_Parameters_v4.xlsx"

        return jsonify({
            "success": True,
            "filename": file.filename,
            "sheets": sheets,
            "sheet_count": len(sheets),
            "default_output": default_output
        })
    except Exception as e:
        return jsonify({"error": f"Failed to inspect Excel file: {str(e)}"}), 500

@app.route("/api/parameter/generate", methods=["POST"])
def api_parameter_generate():
    if "file" not in request.files:
        return jsonify({"error": "No file uploaded."}), 400
    file = request.files["file"]
    if not file.filename:
        return jsonify({"error": "Empty filename."}), 400

    variability_str = request.form.get("variability", "10").strip()
    try:
        variability = float(variability_str)
        if not (0 <= variability <= 30):
            return jsonify({"error": "Variability must be between 0 and 30%."}), 400
    except ValueError:
        return jsonify({"error": "Variability must be a number."}), 400

    seed_str = request.form.get("seed", "").strip()
    if seed_str:
        try:
            seed = int(seed_str)
        except ValueError:
            return jsonify({"error": "Seed must be an integer."}), 400
    else:
        seed = None

    base = Path(file.filename).stem
    default_output_name = f"{base}_Synthetic_Strength_Parameters_v4.xlsx"
    output_name = request.form.get("output_name", "").strip() or default_output_name
    if not output_name.lower().endswith(".xlsx"):
        output_name += ".xlsx"

    safe_out_name = "".join(c for c in output_name if c.isalnum() or c in "._- ")
    if not safe_out_name:
        safe_out_name = default_output_name

    outputs_dir = BASE_DIR / "outputs" / "Parameter"
    outputs_dir.mkdir(parents=True, exist_ok=True)
    out_file_path = outputs_dir / safe_out_name

    with tempfile.NamedTemporaryFile(delete=False, suffix=".xlsx") as tmp:
        temp_input_path = tmp.name
        file.save(temp_input_path)

    try:
        total_processed, processed_sheets, skipped_sheets = param_engine.process_excel(
            temp_input_path,
            str(out_file_path),
            variability,
            seed
        )

        if total_processed == 0:
            return jsonify({
                "error": "No usable data rows with valid Depth and Field SPT N found in any worksheet.",
                "skipped_sheets": skipped_sheets
            }), 400

        preview_wb = load_workbook(str(out_file_path), data_only=True, read_only=True)
        preview_rows = []
        preview_by_sheet = {}

        for sheet_name, p_count in processed_sheets:
            if sheet_name in preview_wb.sheetnames:
                ws = preview_wb[sheet_name]
                sheet_rows = []
                for r_idx in range(10, ws.max_row + 1):
                    depth_val = ws.cell(r_idx, 1).value
                    if depth_val is None or str(depth_val).strip() == "":
                        continue
                    field_n_val = ws.cell(r_idx, 7).value
                    density = ws.cell(r_idx, 47).value
                    if density is None:
                        continue

                    def fmt_num(v, digits=3):
                        try:
                            return f"{float(v):.{digits}f}"
                        except (ValueError, TypeError):
                            return str(v) if v is not None else ""

                    row_dict = {
                        "sheet": sheet_name,
                        "row": r_idx,
                        "depth": fmt_num(depth_val, 2),
                        "field_n": fmt_num(field_n_val, 1),
                        "density": fmt_num(density, 3),
                        "obd": fmt_num(ws.cell(r_idx, 48).value, 3),
                        "pa": fmt_num(ws.cell(r_idx, 49).value, 3),
                        "cn_raw": fmt_num(ws.cell(r_idx, 50).value, 3),
                        "cn": fmt_num(ws.cell(r_idx, 51).value, 3),
                        "n1": fmt_num(ws.cell(r_idx, 52).value, 3),
                        "n1_prime": fmt_num(ws.cell(r_idx, 53).value, 3),
                        "pi": fmt_num(ws.cell(r_idx, 54).value, 3),
                        "li": fmt_num(ws.cell(r_idx, 55).value, 3),
                        "cf": fmt_num(ws.cell(r_idx, 56).value, 3),
                        "soil": str(ws.cell(r_idx, 57).value or ""),
                        "k_pi": fmt_num(ws.cell(r_idx, 58).value, 3),
                        "f_li": fmt_num(ws.cell(r_idx, 59).value, 3),
                        "f_c": fmt_num(ws.cell(r_idx, 60).value, 3),
                        "su": fmt_num(ws.cell(r_idx, 61).value, 3),
                        "ucs": fmt_num(ws.cell(r_idx, 62).value, 3),
                        "c_dr": fmt_num(ws.cell(r_idx, 63).value, 3),
                        "phi_dr": fmt_num(ws.cell(r_idx, 64).value, 2),
                        "cu_uu": fmt_num(ws.cell(r_idx, 65).value, 3),
                        "phi_uu": fmt_num(ws.cell(r_idx, 66).value, 2),
                    }
                    sheet_rows.append(row_dict)
                    preview_rows.append(row_dict)
                preview_by_sheet[sheet_name] = sheet_rows
        preview_wb.close()

        return jsonify({
            "success": True,
            "count": total_processed,
            "processed_sheets": processed_sheets,
            "skipped_sheets": skipped_sheets,
            "excel_filename": safe_out_name,
            "download_url": f"/outputs/Parameter/{safe_out_name}",
            "preview_by_sheet": preview_by_sheet,
            "rows": preview_rows
        })

    except Exception as e:
        return jsonify({"error": f"Processing error: {str(e)}"}), 500
    finally:
        try:
            if os.path.exists(temp_input_path):
                os.remove(temp_input_path)
        except Exception:
            pass

@app.route("/api/parameter/calculate", methods=["POST"])
def api_parameter_calculate():
    data = request.json or {}
    rows = data.get("rows", [])
    variability = float(data.get("variability", 10.0))
    seed_str = str(data.get("random_seed", "")).strip()
    seed = int(seed_str) if seed_str and seed_str.isdigit() else 42
    rng = random.Random(seed)

    if not rows:
        return jsonify({"error": "At least one data row is required."}), 400

    results = []
    for i, r in enumerate(rows, 1):
        try:
            depth = float(r.get("depth", 1.5))
            field_n = float(r.get("field_n", 10.0))
            gravel = float(r.get("gravel", 0.0))
            sand = float(r.get("sand", 0.0))
            silt = float(r.get("silt", 0.0))
            clay = float(r.get("clay", 0.0))
            ll = float(r.get("ll", 0.0))
            pl = float(r.get("pl", 0.0))
            moisture = float(r.get("moisture", 0.0))
        except (ValueError, TypeError):
            return jsonify({"error": f"Row {i}: All physical input values must be numeric."}), 400

        fines = silt + clay
        fine_sand_or_silt = (fines > 0 and clay <= silt and sand >= clay)
        spt = param_engine.calculate_corrected_spt(field_n, depth, fine_sand_or_silt)
        synth = param_engine.synthetic_calculation(
            spt["N1_prime"], gravel, sand, silt, clay, ll, pl, moisture, rng, variability
        )

        results.append({
            "row": i,
            "depth": depth, "field_n": field_n,
            "gravel": gravel, "sand": sand, "silt": silt, "clay": clay,
            "ll": ll, "pl": pl, "moisture": moisture,
            "soil_type": synth[3],
            "density": round(spt["Density_gmcc"], 3),
            "obd": round(spt["Effective_Overburden_kgcm2"], 3),
            "pa": round(spt["Pa_kgcm2"], 3),
            "cn_raw": round(spt["CN_Raw"], 3),
            "cn": round(spt["CN"], 3),
            "n1": round(spt["N1"], 3),
            "n1_prime": round(spt["N1_prime"], 3),
            "pi": round(synth[0], 3),
            "li": round(synth[1], 3),
            "cf": round(synth[2], 3),
            "k_pi": round(synth[4], 3),
            "f_li": round(synth[5], 3),
            "f_c": round(synth[6], 3),
            "su_kgcm2": round(synth[7], 3),
            "ucs_kgcm2": round(synth[8], 3),
            "c_dr_kgcm2": round(synth[9], 3),
            "phi_dr_deg": round(synth[10], 2),
            "cu_uu_kgcm2": round(synth[11], 3),
            "phi_uu_deg": round(synth[12], 2)
        })

    return jsonify({"success": True, "count": len(results), "results": results})

# =========================================================================
# Shrinkage Limit - Exact V5 Observation Generator & Excel Format
# =========================================================================

import shrinkage_limit_generator_v5 as shrink_engine

@app.route("/api/shrinkage/dishes", methods=["GET"])
def api_shrinkage_dishes():
    dishes_list = []
    for d_id in shrink_engine.CONTAINERS:
        d = shrink_engine.DISHES[d_id]
        vol = shrink_engine.dish_volume_cm3(d["diameter"], d["height"])
        dishes_list.append({
            "dish": d_id,
            "weight": d["weight"],
            "diameter": d["diameter"],
            "height": d["height"],
            "volume": round(vol, 2)
        })
    return jsonify({
        "dishes": dishes_list,
        "soil_classes": shrink_engine.SOIL_CLASSES,
        "sample_types": shrink_engine.SAMPLE_TYPES
    })

@app.route("/api/shrinkage/generate", methods=["POST"])
def api_shrinkage_generate():
    data = request.json or {}
    rows = data.get("rows", [])
    bh_id = data.get("bh_id", "BH-01").strip() or "BH-01"

    if not rows:
        return jsonify({"error": "Please enter at least one shrinkage sample."}), 400

    results = []
    for i, r in enumerate(rows, 1):
        try:
            depth = float(r.get("depth", 1.5))
            ll = float(r.get("ll", 40.0))
            pl = float(r.get("pl", 22.0))
        except (ValueError, TypeError):
            return jsonify({"error": f"Row {i}: Depth, LL, and PL must be numeric."}), 400

        sample_type = r.get("sample_type", "UDS")
        soil_class = r.get("soil_class", "CI")
        dish_id = r.get("dish", "SL_01")

        if dish_id not in shrink_engine.DISHES:
            dish_id = "SL_01"

        valid_class, msg = shrink_engine.validate_soil_class(ll, soil_class)
        if not valid_class:
            return jsonify({"error": f"Row {i}: {msg}"}), 400

        dish_data = shrink_engine.DISHES[dish_id]
        obs = shrink_engine.generate_observation(ll, pl, soil_class, dish_data)

        w1_val = obs.get("W1", obs.get("w1", dish_data.get("weight", 0.0)))
        w2_val = obs.get("W2", obs.get("w2", 0.0))
        w3_val = obs.get("W3", obs.get("w3", 0.0))
        m1_val = obs.get("M1", obs.get("m1", 0.0))
        m2_val = obs.get("M2", obs.get("m2", 0.0))
        v1_val = obs.get("V1", obs.get("v1", 0.0))
        mv_val = obs.get("Mercury Mass", obs.get("mv", obs.get("mercury_mass", 0.0)))
        v2_val = obs.get("V2", obs.get("v2", 0.0))
        sl_val = obs.get("SL", obs.get("sl", 0.0))

        results.append({
            "bh_id": bh_id,
            "depth": f"{depth:.2f}",
            "sample_type": sample_type,
            "soil_class": soil_class,
            "dish": dish_id,
            "w1": round(float(w1_val), 2),
            "w2": round(float(w2_val), 2),
            "w3": round(float(w3_val), 2),
            "m1": round(float(m1_val), 2),
            "m2": round(float(m2_val), 2),
            "v1": round(float(v1_val), 2),
            "mv": round(float(mv_val), 2),
            "v2": round(float(v2_val), 2),
            "sl": round(float(sl_val), 2)
        })

    return jsonify({"success": True, "results": results})

@app.route("/api/shrinkage/export_excel", methods=["POST"])
def api_shrinkage_export_excel():
    data = request.json or {}
    results = data.get("results", [])
    bh_id = data.get("bh_id", "Shrinkage_Result").strip() or "Shrinkage_Result"
    if not results:
        return jsonify({"error": "No shrinkage data to export."}), 400

    from openpyxl import Workbook
    from openpyxl.styles import Font, Alignment, Border, Side

    wb = Workbook()
    ws = wb.active
    ws.title = "Shrinkage Limit"

    thin = Side(style="thin", color="000000")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)

    ws.merge_cells("A1:N1")
    ws["A1"] = "SHRINKAGE LIMIT TEST OBSERVATION SHEET"
    ws["A1"].font = Font(name="Segoe UI", size=14, bold=True)
    ws["A1"].alignment = Alignment(horizontal="center", vertical="center")

    headers = [
        "BH ID", "Depth (m)", "Sample Type", "Soil Class", "Dish No.",
        "Empty Container Weight (W1), g", "Container + Wet Soil (W2), g",
        "Container + Dry Soil (W3), g", "Wet Soil Mass (M1), g", "Dry Soil Mass (M2), g",
        "Volume of Wet Pat (V1), cm³", "Mass of Displaced Mercury (Mv), g",
        "Volume of Dry Pat (V2), cm³", "Shrinkage Limit (SL), %"
    ]

    for col_idx, h in enumerate(headers, 1):
        cell = ws.cell(row=3, column=col_idx, value=h)
        cell.font = Font(name="Segoe UI", size=9.5, bold=True)
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = border

    for r_idx, item in enumerate(results, 4):
        vals = [
            item.get("bh_id", bh_id), item.get("depth", ""), item.get("sample_type", ""), item.get("soil_class", ""), item.get("dish", ""),
            item.get("w1", 0.0), item.get("w2", 0.0), item.get("w3", 0.0), f"=G{r_idx}-F{r_idx}", f"=H{r_idx}-F{r_idx}",
            item.get("v1", 0.0), item.get("mv", 0.0), f"=L{r_idx}/13.6", f"=((I{r_idx}-J{r_idx})/J{r_idx} - (K{r_idx}-M{r_idx})/J{r_idx})*100"
        ]
        for c_idx, val in enumerate(vals, 1):
            cell = ws.cell(row=r_idx, column=c_idx, value=val)
            cell.font = Font(name="Segoe UI", size=9)
            cell.alignment = Alignment(horizontal="center", vertical="center")
            cell.border = border

    outputs_base = BASE_DIR / "outputs" / "Shrinkage"
    outputs_base.mkdir(parents=True, exist_ok=True)
    filename = f"Shrinkage_Report_{bh_id}.xlsx"
    out_path = outputs_base / filename
    wb.save(out_path)

    return jsonify({
        "success": True,
        "excel_path": str(out_path),
        "filename": filename,
        "excel_filename": filename,
        "excel_url": f"/outputs/Shrinkage/{filename}",
        "download_url": f"/outputs/Shrinkage/{filename}"
    })

# =========================================================================
# Triaxial (UU Triaxial Multi-Depth)
# =========================================================================

import AI_UU_Triaxial_Data_Generator_MultiDepth_EXCEL_IMPORT_ALL_DEPTHS as triax_engine
from datetime import datetime, timedelta

@app.route("/api/triaxial/preview_plot", methods=["POST"])
def api_triaxial_preview_plot():
    import matplotlib
    matplotlib.use("Agg")
    from matplotlib.figure import Figure
    from matplotlib.patches import Arc
    import io
    import base64

    data = request.json or {}
    bh_id = str(data.get("bh_id", "BH-01")).strip() or "BH-01"
    depth = float(data.get("depth", 5.0))
    dia = float(data.get("diameter", 3.80))
    h = float(data.get("height", 7.64))
    density = float(data.get("density", 1.80))
    sample_type = str(data.get("sample_type", "UDS"))
    soil = str(data.get("soil", "Medium Clay"))
    c_val = float(data.get("c", 0.50))
    phi_val = float(data.get("phi", 10.0))
    p1 = float(data.get("p1", 0.50))
    p2 = float(data.get("p2", 1.00))
    p3 = float(data.get("p3", 1.50))
    seed = int(data.get("random_seed", 20260816))
    interval = float(data.get("interval_mm", 0.10))

    start_str = str(data.get("initial_datetime", "05-08-2026 15:47:00")).strip()
    try:
        start_dt = datetime.strptime(start_str, "%d-%m-%Y %H:%M:%S")
    except Exception:
        start_dt = datetime.now()

    try:
        preview_seed = (
            seed
            + int(round(depth * 1000))
            + int(round(dia * 1000))
            + int(round(h * 1000))
            + int(round(density * 1000))
            + int(round(c_val * 10000))
            + int(round(phi_val * 100))
        )
        local_master = random.Random(preview_seed)
        gap_gen = random.Random(seed)
        pressures = [p1, p2, p3]

        all_obs = []
        summaries = []
        current_time = start_dt

        for test_no, pressure in enumerate(pressures, start=1):
            rng = random.Random(local_master.randint(1, 2_000_000_000))
            rows, summary = triax_engine.generate_test(
                test_no,
                pressure,
                c_val,
                phi_val,
                dia,
                h,
                soil,
                current_time,
                rng,
                pressures,
                interval
            )
            for r in rows:
                r["_depth"] = depth
                r["_sample_type"] = sample_type
                r["_test"] = test_no
            all_obs.extend(rows)
            summaries.append(summary)
            current_time = summary["End Time"] + timedelta(minutes=gap_gen.uniform(5.0, 10.0))

        sigma3 = [float(s["Cell Pressure"]) for s in summaries]
        sigma1 = [float(s["Actual Peak sigma1"]) for s in summaries]
        depth_item = {
            "Depth": depth,
            "Sample Type": sample_type
        }

        fig, fit = triax_engine.build_report_mohr_figure(bh_id, depth_item, summaries, sigma3, sigma1)

        buf = io.BytesIO()
        fig.savefig(buf, format="png", dpi=130, bbox_inches="tight")
        buf.seek(0)
        img_b64 = "data:image/png;base64," + base64.b64encode(buf.read()).decode("utf-8")
        matplotlib.pyplot.close(fig)

        status_text = f"Preview: Depth {depth:.2f} m | Fitted Cu = {fit['Cu']:.2f} kg/cm² | Fitted φ = {fit['phi']:.2f}°"

        formatted_obs = []
        for idx_seq, o in enumerate(all_obs, 1):
            dt_str = o["DateTime"].strftime("%d-%m-%Y %H:%M:%S") if isinstance(o["DateTime"], datetime) else str(o["DateTime"])
            formatted_obs.append({
                "Seq#": idx_seq,
                "depth": depth,
                "test": f"Test {o['_test']}",
                "sample_type": sample_type,
                "DateTime": dt_str,
                "STRAIN": round(float(o["STRAIN"]), 3),
                "LOAD": round(float(o["LOAD"]), 3),
                "_test": o["_test"]
            })

        return jsonify({
            "success": True,
            "image_b64": img_b64,
            "status_text": status_text,
            "fitted_cu": fit["Cu"],
            "fitted_phi": fit["phi"],
            "depth": depth,
            "rows": formatted_obs,
            "total_observations": len(formatted_obs)
        })
    except Exception as e:
        return jsonify({"error": str(e)}), 500

@app.route("/api/triaxial/generate", methods=["POST"])
def api_triaxial_generate():
    data = request.json or {}
    rows = data.get("rows", [])
    bh_id = data.get("bh_id", "BH-01").strip() or "BH-01"
    seed_val = int(data.get("random_seed", 20260816))
    interval = float(data.get("interval_mm", 0.10))

    start_str = str(data.get("initial_datetime", "05-08-2026 15:47:00")).strip()
    try:
        current_dt = datetime.strptime(start_str, "%d-%m-%Y %H:%M:%S")
    except Exception:
        current_dt = datetime.now()

    if not rows:
        return jsonify({"error": "At least one depth is required."}), 400

    master_gen = random.Random(seed_val)
    depth_results = []
    all_combined_rows = []
    global_seq = 1
    test_counter = 0

    for depth_no, r in enumerate(rows, 1):
        try:
            depth = float(r.get("depth", 5.0))
            dia = float(r.get("diameter", 3.80))
            h = float(r.get("height", 7.64))
            density = float(r.get("density", 1.80))
            c_val = float(r.get("c", 0.50))
            phi_val = float(r.get("phi", 10.0))
            p1 = float(r.get("p1", 0.50))
            p2 = float(r.get("p2", 1.00))
            p3 = float(r.get("p3", 1.50))
        except (ValueError, TypeError):
            return jsonify({"error": f"Row {depth_no}: All numeric inputs must be valid numbers."}), 400

        soil = r.get("soil", "Medium Clay")
        sample_type = r.get("sample_type", "UDS")
        pressures = [p1, p2, p3]

        # Exact deterministic preview seed from AI_UU_Triaxial_Data_Generator_MultiDepth_EXCEL_IMPORT_ALL_DEPTHS.py
        preview_seed = (
            seed_val
            + int(round(depth * 1000))
            + int(round(dia * 1000))
            + int(round(h * 1000))
            + int(round(density * 1000))
            + int(round(c_val * 10000))
            + int(round(phi_val * 100))
        )
        local_master = random.Random(preview_seed)

        depth_obs = []
        summaries = []

        for local_n, pressure in enumerate(pressures, start=1):
            test_counter += 1
            rng = random.Random(local_master.randint(1, 2_000_000_000))
            t_rows, summary = triax_engine.generate_test(
                local_n, pressure, c_val, phi_val,
                dia, h, soil,
                current_dt, rng, pressures, interval
            )
            for tr in t_rows:
                tr["_depth_no"] = depth_no
                tr["_depth"] = depth
                tr["_sample_type"] = sample_type
                tr["_soil"] = soil
                tr["_diameter"] = dia
                tr["_height"] = h
                tr["_density"] = density
                tr["_c"] = c_val
                tr["_phi"] = phi_val
                tr["_test_global"] = test_counter
                tr["_test"] = local_n
            summary.update({
                "_depth_no": depth_no, "_depth": depth,
                "_sample_type": sample_type, "_soil": soil,
                "_diameter": dia, "_height": h,
                "_density": density,
                "_c": c_val, "_phi": phi_val,
                "_test_global": test_counter,
            })
            depth_obs.extend(t_rows)
            summaries.append(summary)
            current_dt = summary["End Time"] + timedelta(minutes=master_gen.uniform(5.0, 10.0))

        sigma3 = [float(s["Cell Pressure"]) for s in summaries]
        sigma1 = [float(s["Actual Peak sigma1"]) for s in summaries]
        fit = triax_engine.calculate_final_mohr_coulomb_envelope(sigma3, sigma1)

        # Convert summaries datetime for JSON serialization
        summaries_json = []
        for s in summaries:
            s_copy = dict(s)
            if isinstance(s_copy.get("End Time"), datetime):
                s_copy["End Time"] = s_copy["End Time"].strftime("%d-%m-%Y %H:%M:%S")
            summaries_json.append(s_copy)

        formatted_depth_obs = []
        for o in depth_obs:
            dt_str = o["DateTime"].strftime("%d-%m-%Y %H:%M:%S") if isinstance(o["DateTime"], datetime) else str(o["DateTime"])
            row_dict = {
                "Seq#": global_seq,
                "depth": depth,
                "test": f"Test {o['_test']}",
                "sample_type": sample_type,
                "DateTime": dt_str,
                "STRAIN": round(float(o["STRAIN"]), 3),
                "LOAD": round(float(o["LOAD"]), 3),
                "_test": o["_test"]
            }
            formatted_depth_obs.append(row_dict)
            all_combined_rows.append(row_dict)
            global_seq += 1

        depth_results.append({
            "depth": depth,
            "diameter": dia,
            "height": h,
            "density": density,
            "sample_type": sample_type,
            "soil": soil,
            "c": c_val,
            "phi": phi_val,
            "p1": p1, "p2": p2, "p3": p3,
            "pressures": pressures,
            "fitted_cu": round(fit["Cu"], 3),
            "fitted_phi": round(fit["phi"], 2),
            "sigma3": sigma3,
            "sigma1": sigma1,
            "summaries": summaries_json,
            "rows": formatted_depth_obs,
            "Depth": depth,
            "Sample Type": sample_type,
            "Borehole ID": bh_id
        })

    return jsonify({
        "success": True,
        "bh_id": bh_id,
        "depth_results": depth_results,
        "all_rows": all_combined_rows,
        "observations_preview": all_combined_rows,
        "total_observations": len(all_combined_rows)
    })

@app.route("/api/triaxial/export_plots", methods=["POST"])
def api_triaxial_export_plots():
    import matplotlib
    matplotlib.use("Agg")
    import zipfile

    data = request.json or {}
    depth_results = data.get("depth_results") or data.get("results") or []
    bh_id = data.get("bh_id", "Triaxial_Result").strip() or "Triaxial_Result"

    if not depth_results:
        return jsonify({"error": "No triaxial results to export."}), 400

    outputs_base = BASE_DIR / "outputs" / "Triaxial"
    outputs_base.mkdir(parents=True, exist_ok=True)
    safe_bh = "".join(c if c.isalnum() or c in "-_" else "_" for c in bh_id)
    plot_dir = outputs_base / f"{safe_bh}_Mohr_Plots"
    plot_dir.mkdir(parents=True, exist_ok=True)

    created_files = []
    files_info = []

    for item in depth_results:
        depth = float(item.get("depth", item.get("Depth", 0)))
        sample_type = str(item.get("sample_type", item.get("Sample Type", "UDS")))
        summaries = item.get("summaries", [])
        sigma3 = [float(s["Cell Pressure"]) for s in summaries]
        sigma1 = [float(s["Actual Peak sigma1"]) for s in summaries]

        depth_item = {
            "Depth": depth,
            "Sample Type": sample_type
        }

        fig, fit = triax_engine.build_report_mohr_figure(bh_id, depth_item, summaries, sigma3, sigma1)

        filename = f"{safe_bh}_Depth_{depth:.2f}m_{sample_type}_Mohr_Plot.png"
        out_path = plot_dir / filename
        fig.savefig(out_path, dpi=200, bbox_inches="tight")
        
        # Also save copy in outputs/Triaxial root for direct URL serving
        direct_path = outputs_base / filename
        fig.savefig(direct_path, dpi=200, bbox_inches="tight")

        created_files.append(filename)
        files_info.append({
            "filename": filename,
            "url": f"/outputs/Triaxial/{filename}"
        })

    zip_filename = f"{safe_bh}_Mohr_Plots.zip"
    zip_path = outputs_base / zip_filename
    with zipfile.ZipFile(zip_path, "w", zipfile.ZIP_DEFLATED) as zipf:
        for fn in created_files:
            zipf.write(plot_dir / fn, arcname=fn)

    return jsonify({
        "success": True,
        "count": len(created_files),
        "files": files_info,
        "zip_filename": zip_filename,
        "download_zip_url": f"/outputs/Triaxial/{zip_filename}"
    })

@app.route("/api/triaxial/export_csv", methods=["POST"])
def api_triaxial_export_csv():
    import csv
    data = request.json or {}
    bh_id = data.get("bh_id", "Triaxial_Result").strip() or "Triaxial_Result"
    all_rows = data.get("all_rows", [])

    if not all_rows:
        return jsonify({"error": "No triaxial observation rows to export."}), 400

    outputs_base = BASE_DIR / "outputs" / "Triaxial"
    outputs_base.mkdir(parents=True, exist_ok=True)
    safe_bh = "".join(c if c.isalnum() or c in "-_" else "_" for c in bh_id)
    filename = f"{safe_bh}_Triaxial_Observations.csv"
    out_path = outputs_base / filename

    with open(out_path, "w", newline="", encoding="utf-8-sig") as f:
        w = csv.writer(f)
        w.writerow(["Seq#", "DateTime", "STRAIN (mm)", "LOAD (kg)"])
        for r in all_rows:
            w.writerow([
                r.get("Seq#", ""),
                r.get("DateTime", ""),
                f'{float(r.get("STRAIN", 0)):.3f}',
                f'{float(r.get("LOAD", 0)):.3f}'
            ])

    return jsonify({
        "success": True,
        "filename": filename,
        "csv_url": f"/outputs/Triaxial/{filename}"
    })

@app.route("/api/triaxial/export_xlsm", methods=["POST"])
def api_triaxial_export_xlsm():
    import zipfile
    data = request.json or {}
    bh_id = data.get("bh_id", "Triaxial_Result").strip() or "Triaxial_Result"
    depth_results = data.get("depth_results", [])
    all_rows = data.get("all_rows", [])

    if not depth_results:
        return jsonify({"error": "No triaxial results to populate."}), 400

    template_path = TRIAX_DIR / "TRIAXIAL_TEMPLATE.xlsm"
    outputs_base = BASE_DIR / "outputs" / "Triaxial"
    outputs_base.mkdir(parents=True, exist_ok=True)
    safe_bh = "".join(c if c.isalnum() or c in "-_" else "_" for c in bh_id)

    if not template_path.exists():
        return jsonify({"error": "TRIAXIAL_TEMPLATE.xlsm not found in Triaxial Python Script folder."}), 404

    created_files = []
    for item in depth_results:
        depth = float(item.get("depth", item.get("Depth", 0)))
        dia = float(item.get("diameter", 3.80))
        h = float(item.get("height", 7.64))
        density = float(item.get("density", 1.80))
        pressures = item.get("pressures", [0.5, 1.0, 1.5])
        sample_type = str(item.get("sample_type", item.get("Sample Type", "UDS")))

        # Filter rows for this depth
        d_rows = item.get("rows")
        if not d_rows:
            d_rows = [r for r in all_rows if float(r.get("depth", -1)) == depth]

        filename = f"{safe_bh}_Depth_{depth:.2f}m_{sample_type}.xlsm"
        out_path = outputs_base / filename

        triax_engine.populate_xlsm(
            str(template_path), str(out_path), d_rows,
            dia, h, density, pressures,
            borehole=bh_id, depth=depth
        )
        created_files.append((filename, out_path))

    zip_filename = f"{safe_bh}_Triaxial_Populated_Templates.zip"
    zip_path = outputs_base / zip_filename
    with zipfile.ZipFile(zip_path, "w", zipfile.ZIP_DEFLATED) as zipf:
        for fn, fp in created_files:
            zipf.write(fp, arcname=fn)

    return jsonify({
        "success": True,
        "count": len(created_files),
        "zip_filename": zip_filename,
        "download_zip_url": f"/outputs/Triaxial/{zip_filename}"
    })

@app.route("/api/triaxial/export_all", methods=["POST"])
def api_triaxial_export_all():
    import matplotlib
    matplotlib.use("Agg")
    import zipfile
    import csv

    data = request.json or {}
    bh_id = data.get("bh_id", "Triaxial_Result").strip() or "Triaxial_Result"
    depth_results = data.get("depth_results") or data.get("results") or []
    all_rows = data.get("all_rows", [])

    if not depth_results:
        return jsonify({"error": "No triaxial results to export."}), 400

    template_path = TRIAX_DIR / "TRIAXIAL_TEMPLATE.xlsm"
    if not template_path.exists():
        return jsonify({"error": "TRIAXIAL_TEMPLATE.xlsm not found in Triaxial Python Script folder."}), 404

    outputs_base = BASE_DIR / "outputs" / "Triaxial"
    outputs_base.mkdir(parents=True, exist_ok=True)
    safe_bh = "".join(c if c.isalnum() or c in "-_" else "_" for c in bh_id)

    # 1. Generate Mohr Plot PNGs
    plot_files = []
    for item in depth_results:
        depth = float(item.get("depth", item.get("Depth", 0)))
        sample_type = str(item.get("sample_type", item.get("Sample Type", "UDS")))
        summaries = item.get("summaries", [])
        if not summaries:
            continue
        sigma3 = [float(s["Cell Pressure"]) for s in summaries]
        sigma1 = [float(s["Actual Peak sigma1"]) for s in summaries]

        depth_item = {
            "Depth": depth,
            "Sample Type": sample_type
        }

        fig, fit = triax_engine.build_report_mohr_figure(bh_id, depth_item, summaries, sigma3, sigma1)
        plot_filename = f"{safe_bh}_Depth_{depth:.2f}m_{sample_type}_Mohr_Plot.png"
        plot_path = outputs_base / plot_filename
        fig.savefig(plot_path, dpi=200, bbox_inches="tight")
        matplotlib.pyplot.close(fig)
        plot_files.append((plot_filename, plot_path))

    # 2. Generate CSV Observation file
    csv_filename = f"{safe_bh}_Triaxial_Observations.csv"
    csv_path = outputs_base / csv_filename
    with open(csv_path, "w", newline="", encoding="utf-8-sig") as f:
        w = csv.writer(f)
        w.writerow(["Seq#", "DateTime", "STRAIN (mm)", "LOAD (kg)"])
        for r in all_rows:
            w.writerow([
                r.get("Seq#", ""),
                r.get("DateTime", ""),
                f'{float(r.get("STRAIN", 0)):.3f}',
                f'{float(r.get("LOAD", 0)):.3f}'
            ])

    # 3. Generate Populated XLSM Template files
    xlsm_files = []
    for item in depth_results:
        depth = float(item.get("depth", item.get("Depth", 0)))
        dia = float(item.get("diameter", 3.80))
        h = float(item.get("height", 7.64))
        density = float(item.get("density", 1.80))
        pressures = item.get("pressures", [0.5, 1.0, 1.5])
        sample_type = str(item.get("sample_type", item.get("Sample Type", "UDS")))

        d_rows = item.get("rows")
        if not d_rows:
            d_rows = [r for r in all_rows if float(r.get("depth", -1)) == depth]

        xlsm_filename = f"{safe_bh}_Depth_{depth:.2f}m_{sample_type}.xlsm"
        xlsm_path = outputs_base / xlsm_filename

        triax_engine.populate_xlsm(
            str(template_path), str(xlsm_path), d_rows,
            dia, h, density, pressures,
            borehole=bh_id, depth=depth
        )
        xlsm_files.append((xlsm_filename, xlsm_path))

    # 4. Package all files into a single ZIP
    zip_filename = f"{safe_bh}_Triaxial_All_Export.zip"
    zip_path = outputs_base / zip_filename
    with zipfile.ZipFile(zip_path, "w", zipfile.ZIP_DEFLATED) as zipf:
        # Add Mohr plots
        for fn, fp in plot_files:
            zipf.write(fp, arcname=fn)
        # Add CSV
        zipf.write(csv_path, arcname=csv_filename)
        # Add XLSM templates
        for fn, fp in xlsm_files:
            zipf.write(fp, arcname=fn)

    return jsonify({
        "success": True,
        "count_plots": len(plot_files),
        "count_xlsm": len(xlsm_files),
        "has_csv": True,
        "zip_filename": zip_filename,
        "download_zip_url": f"/outputs/Triaxial/{zip_filename}"
    })

@app.route("/api/triaxial/import_excel", methods=["POST"])
def api_triaxial_import_excel():
    if "file" not in request.files:
        return jsonify({"error": "No file uploaded."}), 400
    file = request.files["file"]
    if not file.filename:
        return jsonify({"error": "Empty filename."}), 400

    try:
        from openpyxl import load_workbook
        wb = load_workbook(file, data_only=True, read_only=True)
        if "Input" not in wb.sheetnames:
            ws = wb.active
        else:
            ws = wb["Input"]

        borehole = ws["B4"].value if ws["B4"].value is not None else "BH-01"
        start_dt_val = ws["B5"].value
        interval = ws["B6"].value if ws["B6"].value is not None else 0.10
        seed = ws["B7"].value if ws["B7"].value is not None else 20260816

        if isinstance(start_dt_val, datetime):
            start_text = start_dt_val.strftime("%d-%m-%Y %H:%M:%S")
        elif start_dt_val:
            start_text = str(start_dt_val).strip()
        else:
            start_text = "05-08-2026 15:47:00"

        rows = []
        for row_num in range(11, ws.max_row + 1):
            vals = [ws.cell(row_num, c).value for c in range(1, 12)]
            if all(v is None or str(v).strip() == "" for v in vals):
                continue
            if vals[0] is None or str(vals[0]).strip() == "":
                continue

            rows.append({
                "depth": str(vals[0]),
                "diameter": str(vals[1] if vals[1] is not None else "3.80"),
                "height": str(vals[2] if vals[2] is not None else "7.64"),
                "density": str(vals[3] if vals[3] is not None else "1.80"),
                "sample_type": str(vals[4] if vals[4] is not None else "UDS"),
                "soil": str(vals[5] if vals[5] is not None else "Medium Clay"),
                "c": str(vals[6] if vals[6] is not None else "0.50"),
                "phi": str(vals[7] if vals[7] is not None else "10.00"),
                "p1": str(vals[8] if vals[8] is not None else "0.50"),
                "p2": str(vals[9] if vals[9] is not None else "1.00"),
                "p3": str(vals[10] if vals[10] is not None else "1.50")
            })

        wb.close()

        if not rows:
            return jsonify({"error": "No depth rows were found in the Excel file."}), 400

        return jsonify({
            "success": True,
            "bh_id": str(borehole).strip(),
            "initial_datetime": start_text,
            "interval_mm": float(interval) if interval else 0.10,
            "random_seed": int(seed) if seed else 20260816,
            "count": len(rows),
            "rows": rows
        })
    except Exception as e:
        return jsonify({"error": f"Failed to parse Excel input: {str(e)}"}), 500

# =========================================================================
# UCS (Unconfined Compressive Strength Multi-Depth)
# =========================================================================

import UCS_Multiple_Depth_Generator_EXCEL_IMPORT as ucs_engine

def render_ucs_plot(specimen, result, output_path=None, as_base64=False, dpi=150):
    import matplotlib
    matplotlib.use("Agg")
    from matplotlib.figure import Figure
    import io
    import base64

    peak = int(result.get("peak_index", 0))
    strains = [float(v) for v in result["strain"]]
    stresses = [float(v) for v in result["stress"]]
    ucs = stresses[peak]
    dia = float(specimen.get("diameter", 3.8))
    area0 = float(result.get("area0", math.pi * dia * dia / 4.0))

    fig = Figure(figsize=(11.69, 8.27), dpi=dpi, facecolor="white")
    
    navy = "#123B8A"
    red = "#E32620"
    black = "#111111"
    light_blue = "#EAF2FA"

    # Page border
    fig.patches.append(
        matplotlib.patches.Rectangle(
            (0.008, 0.008), 0.984, 0.984,
            fill=False, linewidth=1.0,
            transform=fig.transFigure
        )
    )

    # Header
    fig.text(
        0.50, 0.958,
        "UNCONFINED COMPRESSION TEST (UCS)",
        ha="center", va="center",
        fontsize=17, fontweight="bold", color=navy
    )

    fig.lines.append(
        matplotlib.lines.Line2D(
            [0.02, 0.98], [0.925, 0.925],
            transform=fig.transFigure,
            linewidth=0.8, color=black
        )
    )

    # Top left: BH ID, Depth, Sample Type (Column 1)
    for label, value, y in [
        ("BH ID", str(specimen.get("bh_id", "BH-01")), 0.888),
        ("Depth", f"{float(specimen.get('depth', 0)):.2f} m", 0.848),
        ("Sample Type", str(specimen.get("sample") or specimen.get("sample_type", "UDS")), 0.808),
    ]:
        fig.text(0.035, y, label, fontsize=10.5, fontweight="bold")
        fig.text(0.145, y, ":", fontsize=10.5)
        fig.text(0.165, y, str(value), fontsize=10.5, color=navy)

    # Specimen information (Column 2)
    for (label, value), y in zip(
        [
            ("Specimen Diameter (D)", f"{dia:.3f} cm"),
            ("Initial Area (A₀)", f"{area0:.2f} cm²"),
            ("Initial Height (L₀)", f"{float(specimen.get('height', 7.6)):.3f} cm"),
        ],
        [0.888, 0.848, 0.808]
    ):
        fig.text(0.355, y, label, fontsize=10.5, fontweight="bold")
        fig.text(0.535, y, ":", fontsize=10.5)
        fig.text(0.555, y, str(value), fontsize=10.5, color=navy)

    # Bulk Density & Proving Ring Constant (Column 3)
    for (label, value), y in zip(
        [
            ("Bulk Density (γ)", f"{float(specimen.get('density', 1.8)):.2f} g/cc"),
            ("Proving Ring Const.", f"{float(specimen.get('ring_constant', 0.325)):.4f} kg/div"),
        ],
        [0.888, 0.848]
    ):
        fig.text(0.680, y, label, fontsize=10.5, fontweight="bold")
        fig.text(0.845, y, ":", fontsize=10.5)
        fig.text(0.865, y, str(value), fontsize=10.5, color=navy)

    # Main plot
    ax = fig.add_axes([0.080, 0.340, 0.840, 0.340])
    ax.plot(strains, stresses, color=navy, marker="o", markersize=3.2, linewidth=1.4)
    peak_x = strains[peak]
    ax.scatter([peak_x], [ucs], color=red, edgecolor=black, s=45, zorder=5)
    ax.annotate(
        f"UCS (qᵤ) = {ucs:.2f} kg/cm²",
        xy=(peak_x, ucs),
        xytext=(12, 12),
        textcoords="offset points",
        fontsize=9.2,
        fontweight="bold",
        color=red,
        arrowprops=dict(arrowstyle="->", color=red, linewidth=0.9)
    )
    ax.set_title("STRESS – STRAIN CURVE", fontsize=12.5, fontweight="bold", color=navy, pad=8)
    ax.set_xlabel("Axial Strain, ε (%)", fontsize=10.5, fontweight="bold", color=navy)
    ax.set_ylabel("Compressive Stress, σ₀ (kg/cm²)", fontsize=10.5, fontweight="bold", color=navy)
    ax.grid(True, linestyle="--", linewidth=0.55, alpha=0.35)
    ax.set_xlim(0, max(20.0, max(strains) * 1.02))
    ax.set_ylim(0, max(3.0, max(stresses) * 1.18))

    # Results heading
    fig.lines.append(
        matplotlib.lines.Line2D(
            [0.02, 0.98], [0.285, 0.285],
            transform=fig.transFigure,
            linewidth=0.8, color=black
        )
    )
    fig.text(0.50, 0.263, "UCS TEST RESULTS", ha="center", va="center", fontsize=12.5, fontweight="bold", color=navy)

    # Results table
    table_ax = fig.add_axes([0.020, 0.125, 0.960, 0.095])
    table_ax.set_axis_off()
    columns = [
        "Sample Depth\n(m)", "Sample Type", "Specimen Diameter\n(cm)",
        "Initial Height\n(cm)", "Density\n(g/cc)", "Initial Area\n(cm²)", "UCS (qᵤ)\n(kg/cm²)"
    ]
    values = [[
        f"{float(specimen.get('depth', 0)):.2f}",
        str(specimen.get("sample") or specimen.get("sample_type", "UDS")),
        f"{dia:.3f}",
        f"{float(specimen.get('height', 7.6)):.3f}",
        f"{float(specimen.get('density', 1.8)):.2f}",
        f"{area0:.2f}",
        f"{ucs:.2f}"
    ]]
    tbl = table_ax.table(cellText=values, colLabels=columns, cellLoc="center", colLoc="center", loc="center", bbox=[0, 0, 1, 1])
    tbl.auto_set_font_size(False)
    tbl.set_fontsize(8.7)
    for (row, col), cell in tbl.get_celld().items():
        cell.set_edgecolor(black)
        cell.set_linewidth(0.6)
        if row == 0:
            cell.set_facecolor(light_blue)
            cell.set_text_props(fontweight="bold")
        elif col == 6:
            cell.set_text_props(fontweight="bold", color=red)

    # Remarks
    fig.text(0.020, 0.088, "Remarks:", fontsize=9.5, fontweight="bold")
    fig.text(0.020, 0.060, "1. Unconfined Compression Test conducted as per IS 2720 (Part 10) – 1991.", fontsize=8.8)

    if as_base64:
        buf = io.BytesIO()
        fig.savefig(buf, format="png", dpi=dpi, facecolor="white", edgecolor="white")
        buf.seek(0)
        img_b64 = "data:image/png;base64," + base64.b64encode(buf.read()).decode("utf-8")
        return img_b64
    else:
        fig.savefig(output_path, format="png", dpi=dpi, facecolor="white", edgecolor="white")
        return output_path

@app.route("/api/ucs/generate", methods=["POST"])
def api_ucs_generate():
    data = request.json or {}
    rows = data.get("rows", [])
    bh_id = data.get("bh_id", "BH-01").strip() or "BH-01"

    if not rows:
        return jsonify({"error": "At least one depth row is required."}), 400

    results = []
    for idx, r in enumerate(rows, 1):
        try:
            depth = float(r.get("depth", 1.5))
            dia = float(r.get("diameter", 3.80))
            h = float(r.get("height", 7.64))
            density = float(r.get("density", 1.80))
            assigned_ucs = float(r.get("ucs", 0.85))
            ring_const = float(r.get("ring_constant", r.get("ring", 0.325)))
        except (ValueError, TypeError):
            return jsonify({"error": f"Row {idx}: Numerical parameters must be valid numbers."}), 400

        soil = r.get("soil", "Medium Clay")
        sample_type = r.get("sample_type", r.get("sample", "UDS"))
        seed_str = str(r.get("seed", "")).strip()
        seed = int(seed_str) if seed_str and seed_str.isdigit() else 42 + idx

        gen_res = ucs_engine.generate_ucs(assigned_ucs, dia, h, ring_const, soil, seed)

        specimen = {
            "bh_id": bh_id,
            "depth": depth,
            "diameter": dia,
            "height": h,
            "density": density,
            "ucs": assigned_ucs,
            "ring_constant": ring_const,
            "soil": soil,
            "sample": sample_type,
            "sample_type": sample_type,
            "seed": seed
        }

        curve_points = []
        for i in range(32):
            dial = gen_res["dial"][i]
            cm = dial / 1000.0
            def_mm = dial / 1000.0
            strain_val = gen_res["strain"][i]
            pr_val = gen_res["pr"][i]
            stress_val = gen_res["stress"][i]
            load_val = pr_val * ring_const
            corr_area = gen_res["area0"] / (1.0 - (strain_val / 100.0)) if (1.0 - (strain_val / 100.0)) > 0 else 0.0

            curve_points.append({
                "sr": i + 1,
                "dial": dial,
                "cm": f"{cm:.3f}",
                "deformation": def_mm,
                "strain": f"{strain_val:.3f}",
                "strain_pct": round(strain_val, 3),
                "corrected_area": round(corr_area, 3),
                "pr": pr_val,
                "pr_reading": pr_val,
                "load": f"{load_val:.2f}",
                "load_kg": round(load_val, 3),
                "stress": f"{stress_val:.3f}",
                "stress_kgcm2": round(stress_val, 3)
            })

        actual_ucs = gen_res["actual_ucs"]
        peak_idx = gen_res["peak_index"]
        peak_strain = gen_res["strain"][peak_idx]

        results.append({
            "depth": depth,
            "diameter": dia,
            "height": h,
            "density": density,
            "assigned_ucs": assigned_ucs,
            "ring_constant": ring_const,
            "soil": soil,
            "sample": sample_type,
            "sample_type": sample_type,
            "seed": seed,
            "peak_ucs": round(actual_ucs, 3),
            "actual_ucs": actual_ucs,
            "peak_index": peak_idx,
            "peak_strain": round(peak_strain, 3),
            "status": f"Generated {actual_ucs:.3f}",
            "specimen": specimen,
            "result": gen_res,
            "points": curve_points
        })

    return jsonify({
        "success": True,
        "count": len(results),
        "bh_id": bh_id,
        "results": results
    })

@app.route("/api/ucs/preview_plot", methods=["POST"])
def api_ucs_preview_plot():
    data = request.json or {}
    specimen = data.get("specimen", {})
    result = data.get("result", {})
    if not specimen or not result:
        return jsonify({"error": "Specimen and result data required."}), 400
    try:
        img_b64 = render_ucs_plot(specimen, result, as_base64=True, dpi=130)
        return jsonify({"success": True, "image_b64": img_b64})
    except Exception as e:
        return jsonify({"error": str(e)}), 500

@app.route("/api/ucs/export_plot", methods=["POST"])
def api_ucs_export_plot():
    data = request.json or {}
    bh_id = data.get("bh_id", "BH-01").strip() or "BH-01"
    specimen = data.get("specimen", {})
    result = data.get("result", {})
    if not specimen or not result:
        return jsonify({"error": "Specimen and result data required."}), 400

    specimen["bh_id"] = bh_id
    outputs_base = BASE_DIR / "outputs" / "UCS"
    outputs_base.mkdir(parents=True, exist_ok=True)
    depth_val = float(specimen.get("depth", 0))
    sample = specimen.get("sample") or specimen.get("sample_type", "UDS")
    safe_bh = "".join(c if c.isalnum() or c in "-_" else "_" for c in bh_id)
    filename = f"UCS_Report_{safe_bh}_SampleDepth_{depth_val:.2f}m_{sample}.png"
    out_path = outputs_base / filename

    try:
        render_ucs_plot(specimen, result, output_path=str(out_path), dpi=300)
        return jsonify({
            "success": True,
            "filename": filename,
            "url": f"/outputs/UCS/{filename}"
        })
    except Exception as e:
        return jsonify({"error": str(e)}), 500

@app.route("/api/ucs/export_all_plots", methods=["POST"])
def api_ucs_export_all_plots():
    import zipfile
    data = request.json or {}
    bh_id = data.get("bh_id", "BH-01").strip() or "BH-01"
    results = data.get("results", [])
    if not results:
        return jsonify({"error": "No UCS results to export."}), 400

    outputs_base = BASE_DIR / "outputs" / "UCS"
    outputs_base.mkdir(parents=True, exist_ok=True)
    safe_bh = "".join(c if c.isalnum() or c in "-_" else "_" for c in bh_id)
    plot_dir = outputs_base / f"{safe_bh}_UCS_Plots"
    plot_dir.mkdir(parents=True, exist_ok=True)

    created = []
    for item in results:
        specimen = item.get("specimen", {})
        result = item.get("result", {})
        if not specimen or not result:
            continue
        specimen["bh_id"] = bh_id
        depth_val = float(specimen.get("depth", 0))
        sample = specimen.get("sample") or specimen.get("sample_type", "UDS")
        filename = f"UCS_Report_{safe_bh}_SampleDepth_{depth_val:.2f}m_{sample}.png"
        out_path = plot_dir / filename
        render_ucs_plot(specimen, result, output_path=str(out_path), dpi=300)
        created.append(filename)

    zip_filename = f"{safe_bh}_UCS_Report_Plots.zip"
    zip_path = outputs_base / zip_filename
    with zipfile.ZipFile(zip_path, "w", zipfile.ZIP_DEFLATED) as zipf:
        for f in plot_dir.iterdir():
            if f.is_file() and f.name.endswith(".png"):
                zipf.write(f, arcname=f"{plot_dir.name}/{f.name}")

    return jsonify({
        "success": True,
        "count": len(created),
        "filename": zip_filename,
        "zip_url": f"/outputs/UCS/{zip_filename}"
    })

@app.route("/api/ucs/populate_xlsm", methods=["POST"])
def api_ucs_populate_xlsm():
    import zipfile
    data = request.json or {}
    bh_id = data.get("bh_id", "BH-01").strip() or "BH-01"
    results = data.get("results", [])
    if not results:
        return jsonify({"error": "No UCS results to populate."}), 400

    template_path = UCS_DIR / "UCS Template.xlsm"
    outputs_base = BASE_DIR / "outputs" / "UCS"
    outputs_base.mkdir(parents=True, exist_ok=True)
    safe_bh = "".join(c if c.isalnum() or c in "-_" else "_" for c in bh_id)

    if not template_path.exists():
        return jsonify({"error": "UCS Template.xlsm not found in UCS Python Script folder."}), 404

    plot_dir = outputs_base / f"{safe_bh}_UCS_Plots"
    plot_dir.mkdir(parents=True, exist_ok=True)
    for old_f in plot_dir.glob("*.png"):
        try:
            old_f.unlink()
        except Exception:
            pass

    created_xlsm_files = []
    created_plot_files = []

    for item in results:
        specimen = item.get("specimen", {})
        result = item.get("result", {})
        if not specimen or not result:
            continue
        specimen["bh_id"] = bh_id
        depth_val = float(specimen.get("depth", 0))
        sample = specimen.get("sample") or specimen.get("sample_type", "UDS")

        # 1. Populate XLSM Workbook
        xlsm_filename = f"UCS_{safe_bh}_Depth_{depth_val:.2f}m_{sample}.xlsm"
        xlsm_path = outputs_base / xlsm_filename
        ucs_engine.populate_xlsm(str(template_path), str(xlsm_path), specimen, result)
        created_xlsm_files.append((xlsm_filename, xlsm_path))

        # 2. Render and save report plot PNG
        plot_filename = f"UCS_Report_{safe_bh}_SampleDepth_{depth_val:.2f}m_{sample}.png"
        plot_path = plot_dir / plot_filename
        render_ucs_plot(specimen, result, output_path=str(plot_path), dpi=300)

        direct_plot_path = outputs_base / plot_filename
        render_ucs_plot(specimen, result, output_path=str(direct_plot_path), dpi=300)
        created_plot_files.append((plot_filename, plot_path))

    zip_filename = f"UCS_Package_{safe_bh}.zip"
    zip_path = outputs_base / zip_filename
    with zipfile.ZipFile(zip_path, "w", zipfile.ZIP_DEFLATED) as zipf:
        for fn, fp in created_xlsm_files:
            zipf.write(fp, arcname=fn)
        for fn, fp in created_plot_files:
            zipf.write(fp, arcname=f"{safe_bh}_Plots/{fn}")

    return jsonify({
        "success": True,
        "count": len(created_xlsm_files),
        "plot_count": len(created_plot_files),
        "filename": zip_filename,
        "url": f"/outputs/UCS/{zip_filename}",
        "zip_url": f"/outputs/UCS/{zip_filename}"
    })

@app.route("/api/ucs/import_excel", methods=["POST"])
def api_ucs_import_excel():
    if "file" not in request.files:
        return jsonify({"error": "No file uploaded."}), 400
    file = request.files["file"]
    if not file.filename:
        return jsonify({"error": "Empty filename."}), 400

    try:
        from openpyxl import load_workbook
        wb = load_workbook(file, data_only=True, read_only=True)
        ws = wb.active
        headers = [str(c.value).strip() if c.value is not None else "" for c in ws[1]]

        aliases = {
            "Depth": ["Depth", "Depth (m)"],
            "Diameter (cm)": ["Diameter (cm)", "Diameter"],
            "Height (cm)": ["Height (cm)", "Height"],
            "Density (gm/cc)": ["Density (gm/cc)", "Density (g/cc)", "Density"],
            "Assigned UCS (kg/cm²)": ["Assigned UCS (kg/cm²)", "Assigned UCS", "UCS"],
            "PR Constant (kg/div)": ["PR Constant (kg/div)", "PR Constant", "Ring Constant"],
            "Soil Type": ["Soil Type", "Soil"],
            "Sample Type": ["Sample Type", "Sample"],
            "Random Seed": ["Random Seed", "Seed"],
        }

        index_map = {}
        for key, options in aliases.items():
            found = None
            for opt in options:
                if opt in headers:
                    found = headers.index(opt)
                    break
            if found is None:
                return jsonify({"error": f"Missing required column: {key}"}), 400
            index_map[key] = found

        bh_id = "BH-01"
        for opt in ("BH ID", "BH_ID", "Borehole ID"):
            if opt in headers:
                b_idx = headers.index(opt)
                for r in ws.iter_rows(min_row=2, values_only=True):
                    if r[b_idx] is not None and str(r[b_idx]).strip():
                        bh_id = str(r[b_idx]).strip()
                        break
                break

        rows = []
        for r_idx, excel_row in enumerate(ws.iter_rows(min_row=2, values_only=True), 1):
            if not any(v is not None and str(v).strip() != "" for v in excel_row):
                continue
            rows.append({
                "depth": str(excel_row[index_map["Depth"]] if excel_row[index_map["Depth"]] is not None else "1.50"),
                "diameter": str(excel_row[index_map["Diameter (cm)"]] if excel_row[index_map["Diameter (cm)"]] is not None else "3.792"),
                "height": str(excel_row[index_map["Height (cm)"]] if excel_row[index_map["Height (cm)"]] is not None else "7.611"),
                "density": str(excel_row[index_map["Density (gm/cc)"]] if excel_row[index_map["Density (gm/cc)"]] is not None else "1.84"),
                "ucs": str(excel_row[index_map["Assigned UCS (kg/cm²)"]] if excel_row[index_map["Assigned UCS (kg/cm²)"]] is not None else "1.733"),
                "ring_constant": str(excel_row[index_map["PR Constant (kg/div)"]] if excel_row[index_map["PR Constant (kg/div)"]] is not None else "0.325"),
                "soil": str(excel_row[index_map["Soil Type"]] if excel_row[index_map["Soil Type"]] is not None else "Medium Clay"),
                "sample_type": str(excel_row[index_map["Sample Type"]] if excel_row[index_map["Sample Type"]] is not None else "UDS"),
                "seed": int(excel_row[index_map["Random Seed"]]) if excel_row[index_map["Random Seed"]] is not None else (1000 + r_idx),
                "status": "Not generated"
            })
        wb.close()

        if not rows:
            return jsonify({"error": "No input data rows found in Excel file."}), 400

        return jsonify({
            "success": True,
            "bh_id": bh_id,
            "count": len(rows),
            "rows": rows
        })
    except Exception as e:
        return jsonify({"error": f"Failed to import Excel: {str(e)}"}), 500

if __name__ == "__main__":
    port = int(os.environ.get("PORT", 5000))
    use_ssl = "--ssl" in sys.argv
    if use_ssl:
        print(f"Starting server with HTTPS/SSL on https://127.0.0.1:{port}")
        app.run(host="0.0.0.0", port=port, debug=True, ssl_context="adhoc", threaded=False)
    else:
        print(f"Starting server on http://127.0.0.1:{port} and http://localhost:{port}")
        app.run(host="0.0.0.0", port=port, debug=True, threaded=False)

