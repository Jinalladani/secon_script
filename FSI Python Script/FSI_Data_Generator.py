import tkinter as tk
from tkinter import ttk, messagebox, filedialog
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, Border, Side
import random

# ============================================================
# FREE SWELL INDEX (FSI) DATA GENERATOR
#
# GUI:
#   1. One BH ID input
#   2. Multiple depth rows
#   3. Each depth has its own Sample Type
#   4. Each depth has its own Soil Description
#   5. Vk and Vd are integer mL readings (1 mL resolution)
#   6. FSI = ((Vd - Vk) / Vk) * 100
# ============================================================

SOIL_RANGES = {
    "Sand / sandy soil": (2, 10),
    "Silty soil": (5, 15),
    "CL / low plasticity clay": (8, 25),
    "CI / medium plasticity clay": (20, 45),
    "CH / high plasticity clay": (35, 70),
    "Expansive clay": (50, 100),
    "Black cotton soil": (60, 120),
    "Highly expansive / bentonitic soil": (100, 250),
}

SAMPLE_TYPES = ["UDS", "SPT"]
SOIL_OPTIONS = ["Auto / Random Soil Type"] + list(SOIL_RANGES.keys())


def choose_soil(selection):
    if selection == "Auto / Random Soil Type":
        return random.choice(list(SOIL_RANGES.keys()))
    return selection


def generate_fsi_readings(soil):
    """Generate integer Vk/Vd readings that give FSI within soil range."""
    low, high = SOIL_RANGES[soil]

    valid = []

    # 100 mL cylinder, 1 mL resolution.
    # Vk and Vd therefore remain whole numbers.
    for vk in range(5, 21):
        for vd in range(vk + 1, 101):
            fsi = ((vd - vk) / vk) * 100.0
            if low <= fsi <= high:
                valid.append((vk, vd, round(fsi, 2)))

    if not valid:
        # Fallback; should not normally be needed.
        vk = 10
        target = random.uniform(low, high)
        vd = max(vk + 1, round(vk * (1 + target / 100)))
        vd = min(vd, 100)
        fsi = round(((vd - vk) / vk) * 100, 2)
        return vk, vd, fsi

    # Select a random target within the soil range and choose a
    # nearby integer-reading combination.
    target = random.uniform(low, high)
    valid.sort(key=lambda x: abs(x[2] - target))

    pool = valid[:min(30, len(valid))]
    return random.choice(pool)


def export_excel(rows, filepath):
    """Export generated results using the supplied FSI table structure."""
    wb = Workbook()
    ws = wb.active
    ws.title = "FSI"

    thin = Side(style="thin", color="000000")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)

    # Title
    ws.merge_cells("A2:I2")
    ws["A2"] = "Free Swell Index"
    ws["A2"].font = Font(
        name="Arial", size=14, bold=True, underline="single"
    )
    ws["A2"].alignment = Alignment(
        horizontal="center", vertical="center"
    )

    headers = [
        "SR No",
        "BH ID",
        "Sample ID",
        "Sample\nType",
        "Depth",
        "Initial Reading\nin Kerosene\n(After 24 HRS.)",
        "Final Reading in\nDistilled Water\n(After 24 HRS.)",
        "Difference\nin Reading",
        "Free Swell (%)",
    ]

    subheaders = [
        "",
        "",
        "",
        "",
        "(m)",
        "( Vk )",
        "( Vd )",
        "",
        "(( Vd - Vk ) / Vk) * 100",
    ]

    for col, value in enumerate(headers, start=1):
        cell = ws.cell(row=3, column=col, value=value)
        cell.font = Font(name="Arial", size=10, bold=True)
        cell.alignment = Alignment(
            horizontal="center",
            vertical="center",
            wrap_text=True
        )
        cell.border = border

    for col, value in enumerate(subheaders, start=1):
        cell = ws.cell(row=5, column=col, value=value)
        cell.font = Font(name="Arial", size=9, bold=True)
        cell.alignment = Alignment(
            horizontal="center",
            vertical="center",
            wrap_text=True
        )
        cell.border = border

    for r, item in enumerate(rows, start=6):
        values = [
            item["SR No"],
            item["BH ID"],
            item["Sample ID"],
            item["Sample Type"],
            item["Depth"],
            item["Vk"],
            item["Vd"],
            item["Difference"],
            item["FSI"],
        ]

        for c, value in enumerate(values, start=1):
            cell = ws.cell(row=r, column=c, value=value)
            cell.font = Font(name="Arial", size=10)
            cell.alignment = Alignment(
                horizontal="center",
                vertical="center"
            )
            cell.border = border

        # Keep Vk and Vd as 0-decimal readings.
        ws.cell(row=r, column=6).number_format = "0"
        ws.cell(row=r, column=7).number_format = "0"

        # Calculate Difference and FSI in Excel.
        ws.cell(row=r, column=8).value = f"=G{r}-F{r}"
        ws.cell(row=r, column=9).value = f"=((G{r}-F{r})/F{r})*100"
        ws.cell(row=r, column=9).number_format = "0.00"

        ws.cell(row=r, column=5).number_format = "0.00"

    widths = {
        "A": 8,
        "B": 13,
        "C": 14,
        "D": 12,
        "E": 11,
        "F": 23,
        "G": 25,
        "H": 17,
        "I": 22,
    }

    for col, width in widths.items():
        ws.column_dimensions[col].width = width

    ws.row_dimensions[2].height = 26
    ws.row_dimensions[3].height = 68
    ws.row_dimensions[5].height = 32

    ws.freeze_panes = "A6"
    wb.save(filepath)


class FSIApp:
    def __init__(self, root):
        self.root = root
        self.root.title("Free Swell Index (FSI) Data Generator")
        self.root.geometry("1250x820")
        self.root.minsize(1100, 720)

        self.rows = []
        self.depth_rows = []

        self.build_ui()
        self.add_depth_row()

    def build_ui(self):
        main = ttk.Frame(self.root, padding=12)
        main.pack(fill="both", expand=True)

        ttk.Label(
            main,
            text="Free Swell Index (FSI) Data Generator",
            font=("Arial", 18, "bold")
        ).pack(anchor="center", pady=(0, 12))

        # --------------------------------------------------------
        # 1. SINGLE BH ID
        # --------------------------------------------------------
        bh_frame = ttk.LabelFrame(
            main,
            text="Borehole Information",
            padding=10
        )
        bh_frame.pack(fill="x", pady=(0, 10))

        self.bh_var = tk.StringVar(value="RBH-12")

        ttk.Label(
            bh_frame,
            text="BH ID"
        ).grid(row=0, column=0, padx=5, pady=5, sticky="w")

        ttk.Entry(
            bh_frame,
            textvariable=self.bh_var,
            width=25
        ).grid(row=0, column=1, padx=5, pady=5, sticky="w")

        ttk.Label(
            bh_frame,
            text="Single BH ID for all entered sample depths.",
            font=("Arial", 9, "italic")
        ).grid(row=0, column=2, padx=15, pady=5, sticky="w")

        # --------------------------------------------------------
        # 2. MULTI DEPTH INPUT
        # --------------------------------------------------------
        depth_frame = ttk.LabelFrame(
            main,
            text="Multi-Depth Sample Input",
            padding=8
        )
        depth_frame.pack(fill="x", pady=(0, 10))

        # Header row
        header = ttk.Frame(depth_frame)
        header.pack(fill="x", padx=2)

        header.columnconfigure(0, minsize=55)
        header.columnconfigure(1, minsize=180)
        header.columnconfigure(2, minsize=160)
        header.columnconfigure(3, weight=1, minsize=400)
        header.columnconfigure(4, minsize=95)

        for col, text in enumerate(
            ["#", "Sample Depth (m)", "Sample Type",
             "Soil Description", "Action"]
        ):
            ttk.Label(
                header,
                text=text,
                font=("Arial", 10, "bold"),
                anchor="center"
            ).grid(
                row=0,
                column=col,
                sticky="ew",
                padx=5,
                pady=4
            )

        # Scrollable rows
        holder = ttk.Frame(depth_frame)
        holder.pack(fill="x", pady=(2, 0))

        self.depth_canvas = tk.Canvas(
            holder,
            height=205,
            highlightthickness=0
        )

        self.depth_scrollbar = ttk.Scrollbar(
            holder,
            orient="vertical",
            command=self.depth_canvas.yview
        )

        self.depth_container = ttk.Frame(
            self.depth_canvas
        )

        self.depth_window = self.depth_canvas.create_window(
            (0, 0),
            window=self.depth_container,
            anchor="nw"
        )

        self.depth_container.bind(
            "<Configure>",
            lambda event: self.depth_canvas.configure(
                scrollregion=self.depth_canvas.bbox("all")
            )
        )

        self.depth_canvas.bind(
            "<Configure>",
            lambda event: self.depth_canvas.itemconfigure(
                self.depth_window,
                width=event.width
            )
        )

        self.depth_canvas.configure(
            yscrollcommand=self.depth_scrollbar.set
        )

        self.depth_canvas.pack(
            side="left",
            fill="x",
            expand=True
        )

        self.depth_scrollbar.pack(
            side="right",
            fill="y"
        )

        # Buttons below depth table
        depth_buttons = ttk.Frame(depth_frame)
        depth_buttons.pack(fill="x", pady=(8, 2))

        ttk.Button(
            depth_buttons,
            text="+ Add Depth",
            command=self.add_depth_row
        ).pack(side="left", padx=3)

        ttk.Button(
            depth_buttons,
            text="Clear Depths",
            command=self.clear_depths
        ).pack(side="left", padx=3)

        ttk.Label(
            depth_buttons,
            text=(
                "Each depth has its own Sample Type and "
                "Soil Description."
            ),
            font=("Arial", 9, "italic")
        ).pack(side="left", padx=15)

        # --------------------------------------------------------
        # 3. ACTION BUTTONS
        # --------------------------------------------------------
        action_bar = ttk.Frame(main)
        action_bar.pack(fill="x", pady=(0, 10))

        ttk.Button(
            action_bar,
            text="Generate FSI Data",
            command=self.generate
        ).pack(side="left", padx=(0, 7))

        ttk.Button(
            action_bar,
            text="Clear Generated Data",
            command=self.clear_generated
        ).pack(side="left", padx=7)

        ttk.Button(
            action_bar,
            text="Export to Excel",
            command=self.export
        ).pack(side="left", padx=7)

        self.status_var = tk.StringVar(value="Ready")

        ttk.Label(
            action_bar,
            textvariable=self.status_var
        ).pack(side="right")

        # --------------------------------------------------------
        # 4. PREVIEW
        # --------------------------------------------------------
        preview_frame = ttk.LabelFrame(
            main,
            text="Generated Observation Table",
            padding=8
        )
        preview_frame.pack(fill="both", expand=True)

        columns = (
            "SR No",
            "BH ID",
            "Sample ID",
            "Sample Type",
            "Depth",
            "Soil Description",
            "Vk",
            "Vd",
            "Difference",
            "FSI"
        )

        self.tree = ttk.Treeview(
            preview_frame,
            columns=columns,
            show="headings"
        )

        headings = {
            "SR No": "SR No",
            "BH ID": "BH ID",
            "Sample ID": "Sample ID",
            "Sample Type": "Sample Type",
            "Depth": "Depth (m)",
            "Soil Description": "Soil Description",
            "Vk": "Vk (mL)",
            "Vd": "Vd (mL)",
            "Difference": "Difference",
            "FSI": "Free Swell (%)"
        }

        widths = {
            "SR No": 60,
            "BH ID": 90,
            "Sample ID": 100,
            "Sample Type": 95,
            "Depth": 85,
            "Soil Description": 285,
            "Vk": 75,
            "Vd": 75,
            "Difference": 95,
            "FSI": 110
        }

        for col in columns:
            self.tree.heading(col, text=headings[col])
            self.tree.column(
                col,
                width=widths[col],
                anchor="center"
            )

        yscroll = ttk.Scrollbar(
            preview_frame,
            orient="vertical",
            command=self.tree.yview
        )

        xscroll = ttk.Scrollbar(
            preview_frame,
            orient="horizontal",
            command=self.tree.xview
        )

        self.tree.configure(
            yscrollcommand=yscroll.set,
            xscrollcommand=xscroll.set
        )

        self.tree.grid(
            row=0,
            column=0,
            sticky="nsew"
        )

        yscroll.grid(
            row=0,
            column=1,
            sticky="ns"
        )

        xscroll.grid(
            row=1,
            column=0,
            sticky="ew"
        )

        preview_frame.rowconfigure(0, weight=1)
        preview_frame.columnconfigure(0, weight=1)

        ttk.Label(
            main,
            text=(
                "FSI = ((Vd - Vk) / Vk) × 100   |   "
                "Vk and Vd are whole-number mL readings "
                "(1 mL cylinder resolution)"
            ),
            font=("Arial", 9)
        ).pack(anchor="w", pady=(7, 0))

    # ========================================================
    # MULTI-DEPTH ROW MANAGEMENT
    # ========================================================

    def add_depth_row(self):
        number = len(self.depth_rows) + 1

        depth_var = tk.StringVar(
            value=f"{1.50 + (number - 1) * 1.50:.2f}"
        )

        sample_type_var = tk.StringVar(
            value="UDS"
        )

        soil_var = tk.StringVar(
            value="Auto / Random Soil Type"
        )

        row = ttk.Frame(
            self.depth_container
        )

        row.grid(
            row=number - 1,
            column=0,
            sticky="ew",
            pady=2
        )

        row.columnconfigure(3, weight=1)

        number_label = ttk.Label(
            row,
            text=str(number),
            width=5,
            anchor="center"
        )
        number_label.grid(
            row=0,
            column=0,
            padx=5
        )

        ttk.Entry(
            row,
            textvariable=depth_var,
            width=20
        ).grid(
            row=0,
            column=1,
            padx=5
        )

        ttk.Combobox(
            row,
            textvariable=sample_type_var,
            values=SAMPLE_TYPES,
            state="readonly",
            width=17
        ).grid(
            row=0,
            column=2,
            padx=5
        )

        ttk.Combobox(
            row,
            textvariable=soil_var,
            values=SOIL_OPTIONS,
            state="readonly",
            width=43
        ).grid(
            row=0,
            column=3,
            sticky="ew",
            padx=5
        )

        ttk.Button(
            row,
            text="Remove",
            command=lambda r=row: self.remove_depth_row(r)
        ).grid(
            row=0,
            column=4,
            padx=5
        )

        self.depth_rows.append({
            "frame": row,
            "number": number_label,
            "depth": depth_var,
            "sample_type": sample_type_var,
            "soil": soil_var
        })

        self.refresh_depth_rows()

    def remove_depth_row(self, frame):
        for item in self.depth_rows:
            if item["frame"] is frame:
                item["frame"].destroy()
                self.depth_rows.remove(item)
                break

        if not self.depth_rows:
            self.add_depth_row()
        else:
            self.refresh_depth_rows()

    def refresh_depth_rows(self):
        for i, item in enumerate(
            self.depth_rows,
            start=1
        ):
            item["number"].configure(
                text=str(i)
            )
            item["frame"].grid_configure(
                row=i - 1
            )

    def clear_depths(self):
        for item in self.depth_rows:
            item["frame"].destroy()

        self.depth_rows.clear()
        self.add_depth_row()

    # ========================================================
    # INPUT VALIDATION
    # ========================================================

    def read_inputs(self):
        bh_id = self.bh_var.get().strip()

        if not bh_id:
            raise ValueError(
                "BH ID cannot be empty."
            )

        if not self.depth_rows:
            raise ValueError(
                "Add at least one sample depth."
            )

        inputs = []
        used_depths = set()

        for i, item in enumerate(
            self.depth_rows,
            start=1
        ):
            text = item["depth"].get().strip()

            try:
                depth = float(text)
            except ValueError:
                raise ValueError(
                    f"Invalid Sample Depth in row {i}."
                )

            if depth < 0:
                raise ValueError(
                    f"Sample Depth in row {i} cannot be negative."
                )

            if depth in used_depths:
                raise ValueError(
                    f"Duplicate Sample Depth {depth:g} m."
                )

            used_depths.add(depth)

            sample_type = (
                item["sample_type"].get().strip()
            )

            soil = item["soil"].get().strip()

            if sample_type not in SAMPLE_TYPES:
                raise ValueError(
                    f"Select Sample Type in row {i}."
                )

            if soil not in SOIL_OPTIONS:
                raise ValueError(
                    f"Select Soil Description in row {i}."
                )

            inputs.append({
                "depth": depth,
                "sample_type": sample_type,
                "soil": soil
            })

        return bh_id, inputs

    # ========================================================
    # GENERATE
    # ========================================================

    def generate(self):
        try:
            bh_id, inputs = self.read_inputs()

            generated = []

            for sr_no, item in enumerate(
                inputs,
                start=1
            ):
                soil = choose_soil(
                    item["soil"]
                )

                vk, vd, fsi = generate_fsi_readings(
                    soil
                )

                generated.append({
                    "SR No": sr_no,
                    "BH ID": bh_id,
                    "Sample ID": f"FSI-{sr_no:02d}",
                    "Sample Type": item["sample_type"],
                    "Depth": round(item["depth"], 2),
                    "Soil Description": soil,
                    "Vk": int(vk),
                    "Vd": int(vd),
                    "Difference": int(vd - vk),
                    "FSI": round(
                        ((vd - vk) / vk) * 100,
                        2
                    )
                })

            self.rows = generated

            for item in self.tree.get_children():
                self.tree.delete(item)

            for row in self.rows:
                self.tree.insert(
                    "",
                    "end",
                    values=(
                        row["SR No"],
                        row["BH ID"],
                        row["Sample ID"],
                        row["Sample Type"],
                        f'{row["Depth"]:.2f}',
                        row["Soil Description"],
                        row["Vk"],
                        row["Vd"],
                        row["Difference"],
                        f'{row["FSI"]:.2f}'
                    )
                )

            self.status_var.set(
                f"{len(self.rows)} sample(s) generated"
            )

        except Exception as exc:
            messagebox.showerror(
                "Input Error",
                str(exc)
            )

    # ========================================================
    # CLEAR GENERATED RESULTS
    # ========================================================

    def clear_generated(self):
        self.rows.clear()

        for item in self.tree.get_children():
            self.tree.delete(item)

        self.status_var.set(
            "Generated data cleared"
        )

    # ========================================================
    # EXPORT
    # ========================================================

    def export(self):
        if not self.rows:
            messagebox.showwarning(
                "No Data",
                "Generate FSI data before exporting."
            )
            return

        filepath = filedialog.asksaveasfilename(
            title="Export FSI Excel",
            defaultextension=".xlsx",
            filetypes=[
                ("Excel Workbook", "*.xlsx")
            ]
        )

        if not filepath:
            return

        try:
            export_excel(
                self.rows,
                filepath
            )

            self.status_var.set(
                "Excel exported successfully"
            )

            messagebox.showinfo(
                "Export Complete",
                f"FSI data exported to:\n{filepath}"
            )

        except Exception as exc:
            messagebox.showerror(
                "Export Error",
                str(exc)
            )


if __name__ == "__main__":
    root = tk.Tk()
    app = FSIApp(root)
    root.mainloop()
